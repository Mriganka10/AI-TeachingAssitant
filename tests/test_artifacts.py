from pathlib import Path
from zipfile import ZipFile

from docx import Document
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader

from app.agents.llm import LLMService
from app.artifacts.generator import create_artifacts


def test_teaching_artifacts_accept_nested_model_values(tmp_path: Path) -> None:
    result = {
        "title": {"text": "Time Value of Money"},
        "overview": {"summary": "An MBA Finance teaching package."},
        "learning_objectives": [{"text": "Calculate present value."}],
        "lecture_sections": [
            {
                "title": {"value": "Foundations"},
                "content": [{"text": "Compounding and discounting"}],
            }
        ],
        "discussion_questions": {"text": "Why does timing affect value?"},
        "viva_questions": [{"value": "Define discount rate."}],
    }

    paths = create_artifacts(result, agent_type="teaching", output_dir=tmp_path)

    assert {path.suffix for path in paths} == {".json", ".docx", ".pdf", ".pptx"}
    assert all(path.exists() and path.stat().st_size > 0 for path in paths)


def test_teaching_deck_is_compact_and_cleans_model_markdown(tmp_path: Path) -> None:
    long_text = (
        "Supervised learning uses labelled examples while unsupervised learning discovers "
        "structure in unlabelled data. " * 18
    ) + "[scikit-learn](https://scikit-learn.org/stable/user_guide?utm_source=openai)"
    result = {
        "title": "Supervised and Unsupervised Learning in Machine Learning",
        "overview": {
            "course": "Artificial Intelligence",
            "audience": "B.Tech students",
            "duration_minutes": 60,
            "difficulty": "intermediate",
            "lesson_purpose": long_text,
            "prerequisite_knowledge": [long_text] * 4,
            "evidence_base": [long_text] * 4,
            "core_formulas": [long_text] * 5,
        },
        "learning_objectives": [
            {"id": f"LO{i}", "objective": long_text, "bloom_level": "Apply"} for i in range(1, 6)
        ],
        "lecture_sections": [
            {
                "title": f"Section {index}: A deliberately descriptive lecture title",
                "minutes": 12,
                "content": {
                    "explanation": long_text,
                    "key_points": [long_text] * 6,
                    "worked_example": {"problem": long_text, "solution": long_text},
                    "teaching_tip": long_text,
                    "source_refs": ["S1"],
                },
            }
            for index in range(1, 6)
        ],
        "mcqs": [
            {
                "question": long_text,
                "options": [long_text] * 4,
                "answer": "A",
                "explanation": long_text,
            }
            for _ in range(5)
        ],
        "discussion_questions": [long_text] * 5,
        "sources": [
            {
                "title": "User Guide",
                "year": "2026",
                "url": "https://scikit-learn.org/stable/user_guide?utm_source=openai",
            }
        ],
    }

    paths = create_artifacts(result, agent_type="teaching", output_dir=tmp_path)
    presentation = Presentation(next(path for path in paths if path.suffix == ".pptx"))
    slide_text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )

    assert len(presentation.slides) <= 24
    assert "utm_source=openai" not in slide_text
    assert "[scikit-learn](" not in slide_text


def test_research_artifacts_are_all_professor_formats(tmp_path: Path) -> None:
    result = {
        "title": "Responsible AI in Higher Education",
        "executive_summary": {
            "evidence": "Institutions are adopting AI at different rates.",
            "inference": "Governance maturity may explain part of the variation.",
        },
        "themes": [
            {
                "theme": "Governance",
                "synthesis": {
                    "evidence": "Policies emphasize review and transparency.",
                    "inference": "Implementation quality remains uneven.",
                },
                "papers": ["Example et al., 2026"],
            }
        ],
        "research_gaps": [
            {
                "gap": "Limited longitudinal evidence",
                "evidence": "Most studies are cross-sectional.",
                "inference": "Long-term faculty effects remain uncertain.",
            }
        ],
        "research_questions": ["How does governance maturity affect faculty adoption?"],
        "methodology_suggestions": [
            {"suggestion": "Longitudinal mixed methods", "rationale": "Captures change over time."}
        ],
        "sources": [],
    }

    paths = create_artifacts(result, agent_type="research", output_dir=tmp_path)

    assert {path.suffix for path in paths} == {".json", ".docx", ".pdf", ".pptx"}
    assert all(path.exists() and path.stat().st_size > 0 for path in paths)


def test_pdf_and_pptx_keep_the_requested_subject_in_both_agent_workflows(
    tmp_path: Path,
) -> None:
    for agent_type, request, topic in (
        ("teaching", {"topic": "Time Value of Money", "duration_minutes": 60}, "Time Value of Money"),
        ("research", {"research_topic": "Responsible AI in Education"}, "Responsible AI in Education"),
    ):
        result = LLMService._mock(request)
        paths = create_artifacts(
            result,
            agent_type=agent_type,
            output_dir=tmp_path / agent_type,
        )
        pdf = PdfReader(next(path for path in paths if path.suffix == ".pdf"))
        pdf_text = " ".join(page.extract_text() or "" for page in pdf.pages)
        deck = Presentation(next(path for path in paths if path.suffix == ".pptx"))
        slide_text = " ".join(
            shape.text
            for slide in deck.slides
            for shape in slide.shapes
            if hasattr(shape, "text")
        )

        assert topic.casefold() in pdf_text.casefold()
        assert topic.casefold() in slide_text.casefold()


def test_teacher_reference_details_stay_out_of_student_presentation(tmp_path: Path) -> None:
    result = LLMService._mock({"topic": "Cell Structure", "duration_minutes": 60})
    result["lecture_sections"][0]["content"]["teaching_tip"] = "PRIVATE FACILITATION TIP"
    result["case_study"]["teaching_notes"] = "PRIVATE CASE GUIDANCE"
    result["quality_notes"]["review_priorities"] = ["PRIVATE REVIEW PRIORITY"]
    paths = create_artifacts(result, agent_type="teaching", output_dir=tmp_path)
    by_suffix = {path.suffix: path for path in paths}
    docx_text = "\n".join(p.text for p in Document(by_suffix[".docx"]).paragraphs)
    pdf_text = "\n".join(p.extract_text() or "" for p in PdfReader(by_suffix[".pdf"]).pages)
    deck = Presentation(by_suffix[".pptx"])
    slide_text = "\n".join(
        shape.text for slide in deck.slides for shape in slide.shapes if hasattr(shape, "text")
    )
    for teacher_text in (docx_text, pdf_text):
        assert "PRIVATE FACILITATION TIP" in teacher_text
        assert "PRIVATE CASE GUIDANCE" in teacher_text
        assert "PRIVATE REVIEW PRIORITY" in teacher_text
    for private_text in (
        "PRIVATE FACILITATION TIP", "PRIVATE CASE GUIDANCE", "PRIVATE REVIEW PRIORITY",
        "Professor Review Checklist", "Bloom's Taxonomy Mapping",
    ):
        assert private_text not in slide_text
    assert "Cell Structure" in slide_text
    question_slide = next(
        slide for slide in deck.slides
        if any("Knowledge Check 1" == shape.text for shape in slide.shapes if hasattr(shape, "text"))
    )
    question_text = " ".join(shape.text for shape in question_slide.shapes if hasattr(shape, "text"))
    assert "Answer and explanation" not in question_text
    assert any(
        "ANSWER AND EXPLANATION" in shape.text
        for slide in deck.slides for shape in slide.shapes if hasattr(shape, "text")
    )


def test_research_review_notes_remain_in_reference_not_presentation(tmp_path: Path) -> None:
    result = LLMService._mock({"research_topic": "Cell Structure Research"})
    result["quality_notes"]["review_priorities"] = ["PRIVATE RESEARCH REVIEW"]
    paths = create_artifacts(result, agent_type="research", output_dir=tmp_path)
    by_suffix = {path.suffix: path for path in paths}
    docx_text = "\n".join(p.text for p in Document(by_suffix[".docx"]).paragraphs)
    pdf_text = "\n".join(p.extract_text() or "" for p in PdfReader(by_suffix[".pdf"]).pages)
    deck = Presentation(by_suffix[".pptx"])
    slide_text = "\n".join(
        shape.text for slide in deck.slides for shape in slide.shapes if hasattr(shape, "text")
    )
    assert "PRIVATE RESEARCH REVIEW" in docx_text
    assert "PRIVATE RESEARCH REVIEW" in pdf_text
    assert "PRIVATE RESEARCH REVIEW" not in slide_text
    assert "Cell Structure Research" in slide_text


def test_exporters_preserve_math_and_normalize_mcq_labels(tmp_path: Path) -> None:
    result = {
        "title": "Equation and Assessment QA",
        "overview": {
            "course": "Machine Learning",
            "audience": "B.Tech students",
            "duration_minutes": 30,
            "difficulty": "intermediate",
            "lesson_purpose": "Use mathematical notation accurately.",
            "core_formulas": [
                "Prediction: ŷ = f(x)",
                "MSE = (1/n) Σ_i (y_i - ŷ_i)^2",
                "Centroid μ_k ≈ 3.2",
                "Distance = √((a_1 - b_1)^2 + (a_2 - b_2)^2)",
            ],
        },
        "mcqs": [
            {
                "question": "Which statement is correct?",
                "options": ["A. First", "B) Second", "Option C: Third", "D - Fourth"],
                "answer": "A",
                "explanation": "The first option is correct.",
            }
        ],
        "numerical_problems": [
            {
                "problem": "Compute the mean squared error.",
                "given": ["Predictions ŷ = [2, 4]", "True values y = [1, 5]"],
                "solution_steps": [
                    "Errors y - ŷ = [-1, 1].",
                    "MSE = (1/2) Σ_i (y_i - ŷ_i)^2 = 1.",
                    "Distance = √((a_1 - b_1)^2 + (a_2 - b_2)^2).",
                ],
                "final_answer": "MSE = 1.",
            }
        ],
    }

    paths = create_artifacts(result, agent_type="teaching", output_dir=tmp_path)
    by_suffix = {path.suffix: path for path in paths}

    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(by_suffix[".pdf"]).pages)
    docx_text = "\n".join(paragraph.text for paragraph in Document(by_suffix[".docx"]).paragraphs)
    presentation = Presentation(by_suffix[".pptx"])
    pptx_text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )

    for exported_text in (pdf_text, docx_text, pptx_text):
        assert "A. First" in exported_text
        assert "A. A. First" not in exported_text
        assert "B. B) Second" not in exported_text
        assert "ŷ" in exported_text
        assert "Σ" in exported_text
        assert "₁" in exported_text
        assert "₂" in exported_text
        assert "■" not in exported_text
        assert "□" not in exported_text

    with ZipFile(by_suffix[".docx"]) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "Cambria Math" in document_xml


def test_pptx_callouts_fit_content_without_overlapping_sections(tmp_path: Path) -> None:
    result = {
        "title": "Machine Learning Assessment",
        "overview": {
            "course": "Artificial Intelligence",
            "audience": "B.Tech students",
            "duration_minutes": 30,
            "difficulty": "intermediate",
            "lesson_purpose": "Distinguish learning paradigms and evaluate predictions.",
        },
        "lecture_sections": [
            {
                "title": "Supervised and Unsupervised Learning",
                "minutes": 15,
                "content": {
                    "explanation": "Labels determine whether a model learns a prediction target.",
                    "key_points": [
                        "Supervised learning uses labelled examples.",
                        "Unsupervised learning discovers structure without target labels.",
                    ],
                    "worked_example": {
                        "problem": (
                            "Classify these tasks: predicting an exam score, detecting spam, "
                            "grouping customers without predefined types, and compressing image features."
                        ),
                        "solution": (
                            "The first two tasks are supervised. Customer grouping and feature "
                            "compression are unsupervised."
                        ),
                    },
                    "teaching_tip": (
                        "Draw one table with columns x and y and another with only x. Ask students "
                        "which information the algorithm can use during training."
                    ),
                    "common_errors": [
                        "Do not assume an unlabeled dataset contains a known prediction target."
                    ],
                },
            }
        ],
        "mcqs": [
            {
                "question": "Why does a supervised model need a separate test set?",
                "options": [
                    "Estimate generalization",
                    "Fit parameters",
                    "Create labels",
                    "Scale features",
                ],
                "answer": "A",
                "explanation": (
                    "The test set is not used to fit the model. It estimates performance on data "
                    "that the model did not see during training."
                ),
            }
        ],
        "numerical_problems": [
            {
                "problem": (
                    "A classifier is tested on 12 examples. Compare the true and predicted labels "
                    "and compute classification accuracy."
                ),
                "given": [
                    "Total examples n = 12",
                    "A correct prediction means the predicted label equals the true label",
                    "Accuracy = correct predictions / total predictions",
                ],
                "solution_steps": [
                    "Compare each pair of true and predicted labels.",
                    "Count nine correct predictions.",
                    "Divide 9 by 12 to obtain 0.75.",
                    "Convert the result to 75 percent.",
                ],
                "final_answer": "Accuracy = 0.75 = 75%.",
            }
        ],
    }

    paths = create_artifacts(result, agent_type="teaching", output_dir=tmp_path)
    presentation = Presentation(next(path for path in paths if path.suffix == ".pptx"))

    section_slide = next(
        slide
        for slide in presentation.slides
        if any(
            getattr(shape, "text", "") == "Supervised and Unsupervised Learning"
            for shape in slide.shapes
        )
    )
    application = next(
        shape for shape in section_slide.shapes if shape.text.startswith("APPLICATION")
    )
    misconception_slide, misconception = next(
        (candidate_slide, shape)
        for candidate_slide in presentation.slides
        for shape in candidate_slide.shapes
        if getattr(shape, "text", "").startswith("WATCH OUT FOR")
    )
    if misconception_slide == section_slide:
        assert application.top + application.height <= misconception.top
    else:
        assert any(
            getattr(shape, "text", "") == "Supervised and Unsupervised Learning: Common Misconceptions"
            for shape in misconception_slide.shapes
        )
    assert not any(
        "Draw one table with columns x and y" in getattr(shape, "text", "")
        for slide in presentation.slides for shape in slide.shapes
    )
    assert "…" not in application.text

    numerical_slide = next(
        slide
        for slide in presentation.slides
        if any(getattr(shape, "text", "") == "Numerical Problem 1" for shape in slide.shapes)
    )
    problem = next(shape for shape in numerical_slide.shapes if shape.text.startswith("PROBLEM"))
    solution = next(shape for shape in numerical_slide.shapes if shape.text == "Solution")
    final_answer = next(
        shape for shape in numerical_slide.shapes if shape.text.startswith("FINAL ANSWER")
    )
    assert problem.top + problem.height <= solution.top
    assert final_answer.height < Inches(1.5)
