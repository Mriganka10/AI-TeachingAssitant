import json
import re
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime
from html import escape
from pathlib import Path
from typing import Any
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches as PptxInches
from pptx.util import Pt as PptxPt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFError, TTFont
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

NAVY = "12243F"
BLUE = "2557D6"
CYAN = "22A6B6"
PURPLE = "6654C7"
INK = "253247"
MUTED = "667386"
LIGHT = "F3F6FA"
LINE = "DDE4ED"
WHITE = "FFFFFF"
GREEN = "19845F"
GOLD = "9A6B13"

_OPTION_PREFIX = re.compile(r"^\s*(?:option\s+)?[A-Ha-h]\s*[.):\-]\s*", re.IGNORECASE)
_MATH_CHARACTERS = re.compile(
    r"[\u0370-\u03ff\u2070-\u209f\u2200-\u22ff\u00b5\u0176\u0177\u0232\u0233]"
)
_PDF_FONT_REGULAR = "Helvetica"
_PDF_FONT_BOLD = "Helvetica-Bold"
_SUBSCRIPT_TRANSLATION = str.maketrans("0123456789ijn", "₀₁₂₃₄₅₆₇₈₉ᵢⱼₙ")
_SUPERSCRIPT_TRANSLATION = str.maketrans("0123456789+-", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻")


def _plain_text(value: Any, fallback: str = "") -> str:
    if value is None:
        return fallback
    if isinstance(value, str):
        return _clean_generated_text(value.strip())
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, dict):
        for preferred_key in (
            "title",
            "text",
            "value",
            "name",
            "objective",
            "question",
            "summary",
            "description",
            "evidence",
        ):
            if preferred_key in value and value[preferred_key] is not None:
                return _plain_text(value[preferred_key], fallback)
        return "; ".join(
            f"{_humanize(key)}: {_plain_text(item)}" for key, item in value.items()
        )
    if isinstance(value, (list, tuple)):
        return "\n".join(_plain_text(item) for item in value if item is not None)
    return str(value)


def _clean_generated_text(value: str) -> str:
    """Remove model-facing Markdown and tracking parameters from exported prose."""
    value = re.sub(
        r"\[([^\]]+)\]\((https?://[^)]+)\)",
        lambda match: f"{match.group(1)} ({_clean_url(match.group(2))})",
        value,
    )
    value = re.sub(r"(?<!\()https?://[^\s)]+", lambda match: _clean_url(match.group(0)), value)
    value = value.replace("utm_source=openai&", "").replace("?utm_source=openai", "")
    return _format_math_notation(value)


def _format_math_notation(value: str) -> str:
    """Convert common model notation into portable, readable Unicode mathematics."""
    replacements = {
        r"\hat{y}": "ŷ",
        r"\mu": "μ",
        r"\Sigma": "Σ",
        r"\sum": "Σ",
        r"\approx": "≈",
        r"\times": "×",
        r"\leq": "≤",
        r"\geq": "≥",
    }
    for source, replacement in replacements.items():
        value = value.replace(source, replacement)
    value = value.replace("$", "")
    value = re.sub(
        r"_\{([0-9ijn]+)\}|_([0-9ijn]+)",
        lambda match: (match.group(1) or match.group(2)).translate(_SUBSCRIPT_TRANSLATION),
        value,
    )
    value = re.sub(
        r"\^\{([0-9+\-]+)\}|\^([0-9+\-]+)",
        lambda match: (match.group(1) or match.group(2)).translate(_SUPERSCRIPT_TRANSLATION),
        value,
    )
    return value


def _mcq_option_text(value: Any) -> str:
    """Return an option without a model-supplied A./B)/Option C prefix."""
    return _OPTION_PREFIX.sub("", _plain_text(value), count=1).strip()


def _contains_math(text: Any) -> bool:
    return bool(_MATH_CHARACTERS.search(_plain_text(text)))


def _register_pdf_fonts() -> tuple[str, str]:
    """Register an embedded Unicode font so mathematical glyphs never become squares."""
    global _PDF_FONT_REGULAR, _PDF_FONT_BOLD
    if _PDF_FONT_REGULAR != "Helvetica":
        return _PDF_FONT_REGULAR, _PDF_FONT_BOLD

    candidates = [
        (Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/arialbd.ttf")),
        (
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
            Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ),
        (
            Path(pdfmetrics.__file__).resolve().parents[1] / "fonts" / "Vera.ttf",
            Path(pdfmetrics.__file__).resolve().parents[1] / "fonts" / "VeraBd.ttf",
        ),
    ]
    for regular_path, bold_path in candidates:
        if not (regular_path.exists() and bold_path.exists()):
            continue
        try:
            pdfmetrics.registerFont(TTFont("ProfessorSans", str(regular_path)))
            pdfmetrics.registerFont(TTFont("ProfessorSans-Bold", str(bold_path)))
            pdfmetrics.registerFontFamily(
                "ProfessorSans",
                normal="ProfessorSans",
                bold="ProfessorSans-Bold",
                italic="ProfessorSans",
                boldItalic="ProfessorSans-Bold",
            )
        except TTFError:
            _PDF_FONT_REGULAR = "Helvetica"
            _PDF_FONT_BOLD = "Helvetica-Bold"
        else:
            _PDF_FONT_REGULAR = "ProfessorSans"
            _PDF_FONT_BOLD = "ProfessorSans-Bold"
            break
    return _PDF_FONT_REGULAR, _PDF_FONT_BOLD


def _clean_url(value: str) -> str:
    try:
        parts = urlsplit(value)
        query = urlencode(
            [(key, item) for key, item in parse_qsl(parts.query) if not key.lower().startswith("utm_")]
        )
        return urlunsplit((parts.scheme, parts.netloc, parts.path, query, parts.fragment))
    except ValueError:
        return value


def _shorten(value: Any, limit: int) -> str:
    text = _plain_text(value)
    if len(text) <= limit:
        return text
    shortened = text[: max(1, limit - 1)].rsplit(" ", 1)[0].rstrip(" ,;:")
    return (shortened or text[: limit - 1]).rstrip() + "…"


def _as_list(value: Any) -> list[Any]:
    if value is None:
        return []
    return value if isinstance(value, list) else [value]


def _humanize(value: str) -> str:
    return value.replace("_", " ").strip().title()


def _safe_title(value: Any) -> str:
    text = _plain_text(value, "output")
    return "".join(c if c.isalnum() else "-" for c in text).strip("-")[:60] or "output"


def _color(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def _pptx_color(value: str) -> PptxRGBColor:
    return PptxRGBColor.from_string(value)


def create_artifacts(result: dict, *, agent_type: str, output_dir: Path) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    base = _safe_title(result.get("title", agent_type))
    json_path = output_dir / f"{base}.json"
    json_path.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    builders = (
        (_create_docx, output_dir / f"{base}.docx"),
        (_create_pdf, output_dir / f"{base}.pdf"),
        (_create_pptx, output_dir / f"{base}.pptx"),
    )
    # Each exporter writes a different file and has no shared mutable document state. Running
    # them together shortens the post-generation wait without changing any output format.
    with ThreadPoolExecutor(max_workers=len(builders), thread_name_prefix="artifact") as executor:
        futures = [
            executor.submit(builder, result, destination, agent_type)
            for builder, destination in builders
        ]
        generated = [future.result() for future in futures]
    return [json_path, *generated]


# ---------------------------------------------------------------------------
# DOCX - standard_business_brief preset with academic-report named overrides
# ---------------------------------------------------------------------------


def _set_docx_font(run, *, size: float, color: str = INK, bold: bool = False) -> None:
    font_name = "Cambria Math" if _contains_math(run.text) else "Arial"
    run.font.name = font_name
    fonts = run._element.get_or_add_rPr().rFonts
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attribute), font_name)
    run.font.size = Pt(size)
    run.font.color.rgb = _color(color)
    run.bold = bold


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top: int = 90, start: int = 120, bottom: int = 90, end: int = 120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table) -> None:
    table_properties = table._tbl.tblPr
    borders = table_properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9D9D9")


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    _set_docx_font(run, size=9, color=MUTED)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])


def _configure_docx(doc: Document, title: str, agent_type: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = _color(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color, before, after in (
        ("Title", 26, "000000", 0, 8),
        ("Subtitle", 12, "000000", 0, 18),
        ("Heading 1", 17, "000000", 18, 8),
        ("Heading 2", 13.5, "000000", 13, 6),
        ("Heading 3", 11.5, "000000", 9, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = _color(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "PROFESSOR AI - " + (
        "TEACHING PACKAGE" if agent_type == "teaching" else "RESEARCH SYNTHESIS"
    )
    _set_docx_font(header.runs[0], size=8.5, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    _add_page_number(footer)

    props = doc.core_properties
    props.title = title
    props.subject = "AI-assisted faculty draft for professor review"
    props.author = "Professor AI Workspace"
    props.comments = "Generated material must be reviewed by a professor before use."


def _add_docx_cover(doc: Document, result: dict, agent_type: str) -> None:
    title = _plain_text(result.get("title"), "Professor AI Output")
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(22)
    kicker.paragraph_format.space_after = Pt(10)
    run = kicker.add_run(
        "TEACHING & CURRICULUM"
        if agent_type == "teaching"
        else "RESEARCH & LITERATURE SYNTHESIS"
    )
    _set_docx_font(run, size=9, color=CYAN if agent_type == "teaching" else PURPLE, bold=True)

    title_p = doc.add_paragraph(style="Title")
    title_p.add_run(title)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run(
        "Professor review edition - Generated " + datetime.now(UTC).strftime("%d %B %Y")
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(12)
    label = p.add_run("Faculty review note: ")
    _set_docx_font(label, size=9.5, color="000000", bold=True)
    body = p.add_run(
        "This is AI-assisted draft material. Verify facts, citations, calculations, "
        "assessment suitability, and institutional policy before classroom or publication use."
    )
    _set_docx_font(body, size=9.5, color=INK)


def _add_docx_bullet(doc: Document, text: Any, *, level: int = 0) -> None:
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(_plain_text(text))
    _set_docx_font(run, size=10)


def _add_docx_number(doc: Document, text: Any) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(_plain_text(text))
    _set_docx_font(run, size=10)


def _add_docx_labelled(doc: Document, label: str, value: Any) -> None:
    text = _plain_text(value)
    if not text:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(f"{label}: ")
    _set_docx_font(run, size=10.5, color=NAVY, bold=True)
    body = p.add_run(text)
    _set_docx_font(body, size=10.5)


def _add_docx_callout(doc: Document, label: str, value: Any, fill: str = LIGHT) -> None:
    _add_docx_labelled(doc, label, value)


def _add_docx_table(doc: Document, headers: list[str], rows: list[list[Any]], widths: list[float]):
    if not rows:
        return
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _set_table_borders(table)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        header_cells[index].width = Inches(widths[index])
        header_cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _shade_cell(header_cells[index], NAVY)
        _set_cell_margins(header_cells[index], top=100, bottom=100)
        p = header_cells[index].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        _set_docx_font(run, size=9, color=WHITE, bold=True)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].width = Inches(widths[index])
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cells[index], top=90, bottom=90)
            if row_index % 2:
                _shade_cell(cells[index], "F7F9FC")
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(_plain_text(value))
            _set_docx_font(run, size=9.2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def _create_docx(result: dict, path: Path, agent_type: str) -> Path:
    doc = Document()
    title = _plain_text(result.get("title"), "Professor AI Output")
    _configure_docx(doc, title, agent_type)
    _add_docx_cover(doc, result, agent_type)
    if agent_type == "teaching":
        _render_teaching_docx(doc, result)
    else:
        _render_research_docx(doc, result)
    doc.save(path)
    return path


def _render_teaching_docx(doc: Document, result: dict) -> None:
    overview = result.get("overview") if isinstance(result.get("overview"), dict) else {}
    doc.add_heading("Session at a Glance", level=1)
    _add_docx_table(
        doc,
        ["Course", "Audience", "Duration", "Level"],
        [[
            overview.get("course", "General"),
            overview.get("audience", "University students"),
            f"{overview.get('duration_minutes', '')} minutes",
            overview.get("difficulty", "Intermediate"),
        ]],
        [1.7, 2.0, 1.3, 1.7],
    )
    _add_docx_callout(doc, "Lesson purpose", overview.get("lesson_purpose"))

    prerequisites = _as_list(overview.get("prerequisite_knowledge"))
    if prerequisites:
        doc.add_heading("Prerequisite Knowledge", level=2)
        for item in prerequisites:
            _add_docx_bullet(doc, item)

    formulas = _as_list(overview.get("core_formulas"))
    if formulas:
        doc.add_heading("Core Formula Sheet", level=2)
        _add_docx_table(
            doc,
            ["Concept and Formula"],
            [[item] for item in formulas],
            [6.7],
        )

    objectives = _as_list(result.get("learning_objectives"))
    if objectives:
        doc.add_heading("Learning Objectives", level=1)
        for item in objectives:
            if isinstance(item, dict):
                objective = item.get("objective") or item.get("text") or item
                prefix = f"{item.get('id')}: " if item.get("id") else ""
                _add_docx_number(doc, prefix + _plain_text(objective))
            else:
                _add_docx_number(doc, item)

    sections = _as_list(result.get("lecture_sections"))
    if sections:
        doc.add_heading("60-Minute Session Plan", level=1)
        plan_rows = []
        for section in sections:
            if not isinstance(section, dict):
                plan_rows.append(["", section, ""])
                continue
            content = section.get("content", {})
            focus = ""
            if isinstance(content, dict):
                focus = _plain_text(
                    content.get("key_points")
                    or content.get("classroom_activity")
                    or content.get("worked_example")
                )
            else:
                focus = _plain_text(content)
            plan_rows.append([
                f"{section.get('minutes', '')} min",
                section.get("title", "Lecture Section"),
                focus[:260],
            ])
        _add_docx_table(doc, ["Time", "Section", "Teaching Focus"], plan_rows, [0.8, 2.35, 3.55])

        doc.add_page_break()
        doc.add_heading("Detailed Lecture Notes", level=1)
        for section in sections:
            if not isinstance(section, dict):
                doc.add_heading("Lecture Section", level=2)
                _add_docx_bullet(doc, section)
                continue
            minutes = section.get("minutes")
            suffix = f"  |  {minutes} minutes" if minutes else ""
            doc.add_heading(_plain_text(section.get("title"), "Lecture Section") + suffix, level=2)
            _render_teaching_section_docx(doc, section.get("content", {}))

    _render_teaching_assessment_docx(doc, result)

    case_study = result.get("case_study")
    if isinstance(case_study, dict) and case_study:
        doc.add_page_break()
        doc.add_heading("Case Study", level=1)
        doc.add_heading(_plain_text(case_study.get("title"), "Applied Case"), level=2)
        _add_docx_callout(doc, "Scenario", case_study.get("scenario"), "EEF3FA")
        questions = _as_list(case_study.get("questions"))
        if questions:
            doc.add_heading("Case Questions", level=3)
            for question in questions:
                _add_docx_number(doc, question)
        _add_docx_callout(doc, "Teaching notes", case_study.get("teaching_notes"), "F3F0FB")

    _render_sources_docx(doc, result)
    _render_quality_notes_docx(doc, result)


def _render_teaching_section_docx(doc: Document, content: Any) -> None:
    if not isinstance(content, dict):
        for item in _as_list(content):
            _add_docx_bullet(doc, item)
        return
    _add_docx_labelled(doc, "Explanation", content.get("explanation"))
    key_points = _as_list(content.get("key_points"))
    if key_points:
        doc.add_heading("Key Points", level=3)
        for point in key_points:
            _add_docx_bullet(doc, point)
    for label, key, fill in (
        ("Classroom activity", "classroom_activity", "EAF7F8"),
        ("Concept check", "concept_check", "EEF3FA"),
        ("Mini discussion", "mini_discussion", "F3F0FB"),
        ("Teaching tip", "teaching_tip", "FFF7E7"),
        ("Closing prompt", "closing_prompt", "EEF3FA"),
        ("Expected answer", "expected_answer", "EAF6F0"),
    ):
        _add_docx_callout(doc, label, content.get(key), fill)
    for key in ("worked_example", "worked_examples"):
        examples = _as_list(content.get(key))
        for index, example in enumerate(examples, start=1):
            if isinstance(example, dict):
                doc.add_heading(f"Worked Example {index}", level=3)
                _add_docx_labelled(doc, "Problem", example.get("problem"))
                _add_docx_callout(doc, "Solution", example.get("solution"), "EAF6F0")
            elif example:
                _add_docx_callout(doc, "Worked example", example)
    board_work = _as_list(content.get("board_work"))
    if board_work:
        doc.add_heading("Board / Slide Work", level=3)
        for item in board_work:
            _add_docx_bullet(doc, item)
    common_errors = _as_list(content.get("common_errors"))
    if common_errors:
        doc.add_heading("Common Errors to Address", level=3)
        for item in common_errors:
            _add_docx_bullet(doc, item)
    source_refs = _as_list(content.get("source_refs"))
    if source_refs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run("Source references: " + "; ".join(_plain_text(item) for item in source_refs))
        _set_docx_font(run, size=8.5, color=MUTED)


def _render_teaching_assessment_docx(doc: Document, result: dict) -> None:
    mcqs = _as_list(result.get("mcqs"))
    assignments = _as_list(result.get("assignments"))
    numerical = _as_list(result.get("numerical_problems"))
    bloom = _as_list(result.get("bloom_mapping"))
    discussion = _as_list(result.get("discussion_questions"))
    viva = _as_list(result.get("viva_questions"))
    if any((mcqs, assignments, numerical, bloom, discussion, viva)):
        doc.add_page_break()
        doc.add_heading("Assessment and Engagement Toolkit", level=1)

    if mcqs:
        doc.add_heading("Multiple-Choice Questions", level=2)
        for index, item in enumerate(mcqs, start=1):
            if not isinstance(item, dict):
                _add_docx_number(doc, item)
                continue
            q = doc.add_paragraph()
            q.paragraph_format.space_before = Pt(7)
            q.paragraph_format.space_after = Pt(3)
            run = q.add_run(f"{index}. {_plain_text(item.get('question'))}")
            _set_docx_font(run, size=10.5, color=NAVY, bold=True)
            for option_index, option in enumerate(_as_list(item.get("options"))):
                _add_docx_bullet(
                    doc,
                    f"{chr(65 + option_index)}. {_mcq_option_text(option)}",
                    level=1,
                )
            _add_docx_callout(
                doc,
                "Answer and rationale",
                f"{_plain_text(item.get('answer'))}\n{_plain_text(item.get('explanation'))}",
                "EAF6F0",
            )

    if numerical:
        doc.add_heading("Numerical Problems", level=2)
        for index, item in enumerate(numerical, start=1):
            if isinstance(item, dict):
                doc.add_heading(f"Problem {index}", level=3)
                _add_docx_labelled(doc, "Question", item.get("problem"))
                if item.get("given"):
                    _add_docx_labelled(doc, "Given", item.get("given"))
                steps = item.get("solution_steps") or item.get("solution")
                _add_docx_labelled(doc, "Solution steps", steps)
                _add_docx_labelled(doc, "Final answer", item.get("final_answer"))
            else:
                _add_docx_number(doc, item)

    if assignments:
        doc.add_heading("Assignments", level=2)
        for item in assignments:
            if isinstance(item, dict):
                doc.add_heading(_plain_text(item.get("title"), "Assignment"), level=3)
                _add_docx_callout(doc, "Brief", item.get("prompt"), "EEF3FA")
                deliverables = _as_list(item.get("deliverables"))
                if deliverables:
                    doc.add_paragraph("Deliverables", style="Heading 3")
                    for deliverable in deliverables:
                        _add_docx_bullet(doc, deliverable)
                rubric = _as_list(item.get("rubric"))
                if rubric:
                    rows = []
                    for criterion in rubric:
                        if isinstance(criterion, dict):
                            rows.append([
                                criterion.get("criterion"),
                                f"{criterion.get('weight_percent', '')}%",
                                criterion.get("description"),
                            ])
                        else:
                            rows.append([criterion, "", ""])
                    doc.add_paragraph("Assessment criteria", style="Heading 3")
                    _add_docx_table(
                        doc,
                        ["Criterion", "Weight", "Performance description"],
                        rows,
                        [1.6, 0.8, 4.3],
                    )
            else:
                _add_docx_bullet(doc, item)

    if bloom:
        rows = []
        for item in bloom:
            if isinstance(item, dict):
                rows.append([item.get("objective"), item.get("level"), item.get("assessment")])
            else:
                rows.append([item, "", ""])
        doc.add_heading("Bloom's Taxonomy Mapping", level=2)
        _add_docx_table(
            doc,
            ["Learning Objective", "Bloom Level", "Assessment"],
            rows,
            [2.5, 1.2, 3.0],
        )

    if discussion:
        doc.add_heading("Discussion Questions", level=2)
        for item in discussion:
            _add_docx_number(doc, item)

    if viva:
        doc.add_heading("Viva Questions and Expected Answers", level=2)
        rows = []
        for item in viva:
            if isinstance(item, dict):
                rows.append([item.get("question"), item.get("expected_answer")])
            else:
                rows.append([item, ""])
        _add_docx_table(doc, ["Question", "Expected Answer / Talking Points"], rows, [3.0, 3.7])


def _render_research_docx(doc: Document, result: dict) -> None:
    scope = result.get("scope") if isinstance(result.get("scope"), dict) else {}
    if scope:
        doc.add_heading("Research Scope", level=1)
        _add_docx_labelled(doc, "Discipline", scope.get("discipline"))
        _add_docx_labelled(doc, "Problem statement", scope.get("problem_statement"))
        _add_docx_labelled(doc, "Included", scope.get("inclusion_boundaries"))
        _add_docx_labelled(doc, "Excluded", scope.get("exclusion_boundaries"))
    summary = result.get("executive_summary")
    doc.add_heading("Executive Summary", level=1)
    if isinstance(summary, dict):
        _add_docx_callout(doc, "Evidence", summary.get("evidence"), "EEF3FA")
        _add_docx_callout(doc, "Interpretation", summary.get("inference"), "F3F0FB")
        _add_docx_labelled(doc, "Limitations", summary.get("limitations"))
    else:
        doc.add_paragraph(_plain_text(summary))

    themes = _as_list(result.get("themes"))
    if themes:
        doc.add_heading("Thematic Literature Synthesis", level=1)
        for index, item in enumerate(themes, start=1):
            if not isinstance(item, dict):
                doc.add_heading(f"Theme {index}", level=2)
                doc.add_paragraph(_plain_text(item))
                continue
            doc.add_heading(f"{index}. {_plain_text(item.get('theme'), 'Theme')}", level=2)
            evidence_items = _as_list(item.get("evidence"))
            if evidence_items:
                doc.add_heading("Evidence", level=3)
                for evidence in evidence_items:
                    _add_docx_bullet(doc, evidence)
            synthesis = item.get("synthesis")
            if isinstance(synthesis, dict):
                _add_docx_labelled(doc, "Evidence", synthesis.get("evidence"))
                _add_docx_labelled(doc, "Interpretation", synthesis.get("inference"))
            else:
                doc.add_paragraph(_plain_text(synthesis))
            papers = _as_list(item.get("papers"))
            if papers:
                p = doc.add_paragraph()
                run = p.add_run("Studies represented: " + "; ".join(_plain_text(x) for x in papers))
                _set_docx_font(run, size=8.5, color=MUTED)

    methods = _as_list(result.get("methodology_comparison"))
    if methods:
        doc.add_page_break()
        doc.add_heading("Methodology Comparison", level=1)
        for index, item in enumerate(methods, start=1):
            if not isinstance(item, dict):
                _add_docx_number(doc, item)
                continue
            doc.add_heading(f"{index}. {_plain_text(item.get('method'), 'Method')}", level=2)
            _add_docx_labelled(doc, "Strengths", item.get("strengths"))
            _add_docx_labelled(doc, "Limitations", item.get("limitations"))
            _add_docx_labelled(doc, "Suitability", item.get("suitability"))
            papers = _as_list(item.get("studies") or item.get("papers"))
            if papers:
                _add_docx_labelled(doc, "Representative studies", "; ".join(map(_plain_text, papers)))

    gaps = _as_list(result.get("research_gaps"))
    if gaps:
        doc.add_heading("Research Gaps", level=1)
        for index, item in enumerate(gaps, start=1):
            if isinstance(item, dict):
                doc.add_heading(f"{index}. {_plain_text(item.get('gap'), 'Research gap')}", level=2)
                _add_docx_callout(doc, "Evidence for the gap", item.get("evidence"), "EEF3FA")
                _add_docx_callout(doc, "Research implication", item.get("inference"), "F3F0FB")
                _add_docx_labelled(doc, "Confidence", item.get("confidence"))
            else:
                _add_docx_number(doc, item)

    questions = _as_list(result.get("research_questions"))
    if questions:
        doc.add_heading("Proposed Research Questions", level=1)
        for question in questions:
            if isinstance(question, dict):
                _add_docx_number(doc, question.get("question"))
                _add_docx_labelled(doc, "Rationale", question.get("rationale"))
                _add_docx_labelled(
                    doc, "Key variables or constructs", question.get("key_variables_or_constructs")
                )
            else:
                _add_docx_number(doc, question)

    future = _as_list(result.get("future_scope"))
    if future:
        doc.add_heading("Future Scope", level=1)
        rows = []
        for item in future:
            if isinstance(item, dict):
                rows.append([item.get("area"), item.get("scope")])
            else:
                rows.append([item, ""])
        _add_docx_table(doc, ["Research Area", "Opportunity"], rows, [2.0, 4.7])

    suggestions = _as_list(result.get("methodology_suggestions"))
    if suggestions:
        doc.add_heading("Methodology Recommendations", level=1)
        for index, item in enumerate(suggestions, start=1):
            if isinstance(item, dict):
                doc.add_heading(
                    f"{index}. {_plain_text(item.get('suggestion'), 'Recommendation')}", level=2
                )
                _add_docx_labelled(doc, "Rationale", item.get("rationale"))
                _add_docx_labelled(doc, "Data collection", item.get("data_collection"))
                _add_docx_labelled(doc, "Analysis", item.get("analysis"))
            else:
                _add_docx_number(doc, item)

    references = _as_list(result.get("apa_references"))
    if references:
        doc.add_page_break()
        doc.add_heading("APA References", level=1)
        for reference in references:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.space_after = Pt(5)
            p.add_run(_plain_text(reference))

    _render_sources_docx(doc, result)
    _render_quality_notes_docx(doc, result)


def _render_sources_docx(doc: Document, result: dict) -> None:
    sources = _as_list(result.get("sources"))
    evidence = result.get("overview", {}).get("evidence_base", []) if isinstance(
        result.get("overview"), dict
    ) else []
    if not sources and not evidence:
        return
    doc.add_page_break()
    doc.add_heading("Sources and Further Reading", level=1)
    for item in sources:
        if isinstance(item, dict):
            title = _plain_text(item.get("title"), "Source")
            year = _plain_text(item.get("year"))
            url = _plain_text(item.get("url"))
            text = title + (f" ({year})" if year else "") + (f"\n{url}" if url else "")
        else:
            text = _plain_text(item)
        _add_docx_number(doc, text)
    if evidence:
        doc.add_heading("Evidence Base Named in the Teaching Plan", level=2)
        for item in _as_list(evidence):
            _add_docx_bullet(doc, item)


def _render_quality_notes_docx(doc: Document, result: dict) -> None:
    notes = result.get("quality_notes")
    if not isinstance(notes, dict):
        return
    doc.add_heading("Professor Review Checklist", level=1)
    for heading, key in (
        ("Limitations", "limitations"),
        ("Unverified Claims", "unverified_claims"),
        ("Review Priorities", "review_priorities"),
    ):
        values = _as_list(notes.get(key))
        if values:
            doc.add_heading(heading, level=2)
            for value in values:
                _add_docx_bullet(doc, value)


# ---------------------------------------------------------------------------
# PDF - academic report layout mirroring the DOCX semantic structure
# ---------------------------------------------------------------------------


def _pdf_styles():
    regular_font, bold_font = _register_pdf_fonts()
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "AcademicTitle",
            parent=base["Title"],
            fontName=bold_font,
            fontSize=23,
            leading=28,
            textColor=colors.HexColor(f"#{NAVY}"),
            alignment=TA_LEFT,
            spaceAfter=12,
        ),
        "subtitle": ParagraphStyle(
            "AcademicSubtitle",
            parent=base["Normal"],
            fontName=regular_font,
            fontSize=9,
            leading=12,
            textColor=colors.HexColor(f"#{MUTED}"),
            spaceAfter=16,
        ),
        "h1": ParagraphStyle(
            "AcademicH1",
            parent=base["Heading1"],
            fontName=bold_font,
            fontSize=15,
            leading=18,
            textColor=colors.black,
            spaceBefore=14,
            spaceAfter=7,
        ),
        "h2": ParagraphStyle(
            "AcademicH2",
            parent=base["Heading2"],
            fontName=bold_font,
            fontSize=11.5,
            leading=14,
            textColor=colors.black,
            spaceBefore=10,
            spaceAfter=5,
        ),
        "h3": ParagraphStyle(
            "AcademicH3",
            parent=base["Heading3"],
            fontName=bold_font,
            fontSize=10,
            leading=12,
            textColor=colors.black,
            spaceBefore=7,
            spaceAfter=4,
        ),
        "body": ParagraphStyle(
            "AcademicBody",
            parent=base["BodyText"],
            fontName=regular_font,
            fontSize=9.2,
            leading=12.5,
            textColor=colors.HexColor(f"#{INK}"),
            spaceAfter=5,
        ),
        "small": ParagraphStyle(
            "AcademicSmall",
            parent=base["BodyText"],
            fontName=regular_font,
            fontSize=7.8,
            leading=10,
            textColor=colors.HexColor(f"#{MUTED}"),
        ),
        "bullet": ParagraphStyle(
            "AcademicBullet",
            parent=base["BodyText"],
            fontName=regular_font,
            fontSize=9.2,
            leading=12.5,
            leftIndent=14,
            firstLineIndent=-7,
            bulletIndent=3,
            spaceAfter=3,
        ),
        "callout": ParagraphStyle(
            "AcademicCallout",
            parent=base["BodyText"],
            fontName=regular_font,
            fontSize=9,
            leading=12,
            textColor=colors.HexColor(f"#{INK}"),
        ),
    }


def _pdf_paragraph(text: Any, style) -> Paragraph:
    return Paragraph(escape(_plain_text(text)).replace("\n", "<br/>"), style)


def _pdf_bullets(story: list, values: Any, styles) -> None:
    for item in _as_list(values):
        story.append(Paragraph("• " + escape(_plain_text(item)), styles["bullet"]))


def _pdf_callout(story: list, label: str, value: Any, styles, fill: str = "F3F6FA"):
    text = _plain_text(value)
    if not text:
        return
    story.append(
        Paragraph(
            f"<b>{escape(label)}:</b> {escape(text).replace(chr(10), '<br/>')}",
            styles["body"],
        )
    )


def _pdf_table(story: list, headers: list[str], rows: list[list[Any]], widths: list[float], styles):
    if not rows:
        return
    data = [[_pdf_paragraph(h, styles["small"]) for h in headers]]
    for row in rows:
        data.append([_pdf_paragraph(value, styles["small"]) for value in row])
    table = Table(data, colWidths=[width * inch for width in widths], repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), _PDF_FONT_BOLD),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor(f"#{LINE}")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
            ]
        )
    )
    story.extend([table, Spacer(1, 7)])


def _pdf_page(canvas, doc, agent_type: str):
    canvas.saveState()
    width, height = LETTER
    canvas.setFont(_PDF_FONT_BOLD, 7.5)
    canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
    label = "PROFESSOR AI - " + (
        "TEACHING PACKAGE" if agent_type == "teaching" else "RESEARCH SYNTHESIS"
    )
    canvas.drawString(0.82 * inch, height - 0.35 * inch, label)
    canvas.setFont(_PDF_FONT_REGULAR, 8)
    canvas.drawRightString(width - 0.82 * inch, 0.38 * inch, f"Page {doc.page}")
    canvas.restoreState()


def _create_pdf(result: dict, path: Path, agent_type: str) -> Path:
    title = _plain_text(result.get("title"), "Professor AI Output")
    styles = _pdf_styles()
    story = [
        Spacer(1, 0.2 * inch),
        _pdf_paragraph(title, styles["title"]),
        _pdf_paragraph(
            "Professor review edition - AI-assisted draft - "
            + datetime.now(UTC).strftime("%d %B %Y"),
            styles["subtitle"],
        ),
    ]
    _pdf_callout(
        story,
        "Faculty review note",
        "Verify facts, citations, calculations, assessment suitability, and institutional "
        "policy before classroom or publication use.",
        styles,
        "EAF2FB",
    )
    if agent_type == "teaching":
        _render_teaching_pdf(story, result, styles)
    else:
        _render_research_pdf(story, result, styles)
    document = SimpleDocTemplate(
        str(path),
        pagesize=LETTER,
        rightMargin=0.82 * inch,
        leftMargin=0.82 * inch,
        topMargin=0.64 * inch,
        bottomMargin=0.62 * inch,
        title=title,
        author="Professor AI Workspace",
        subject="AI-assisted faculty draft",
    )
    document.build(
        story,
        onFirstPage=lambda canvas, doc: _pdf_page(canvas, doc, agent_type),
        onLaterPages=lambda canvas, doc: _pdf_page(canvas, doc, agent_type),
    )
    return path


def _render_teaching_pdf(story: list, result: dict, styles) -> None:
    overview = result.get("overview") if isinstance(result.get("overview"), dict) else {}
    story.append(_pdf_paragraph("Session at a Glance", styles["h1"]))
    _pdf_table(
        story,
        ["Course", "Audience", "Duration", "Level"],
        [[overview.get("course"), overview.get("audience"),
          f"{overview.get('duration_minutes', '')} minutes", overview.get("difficulty")]],
        [1.6, 2.0, 1.35, 1.75],
        styles,
    )
    _pdf_callout(story, "Lesson purpose", overview.get("lesson_purpose"), styles, "EEF3FA")
    if overview.get("prerequisite_knowledge"):
        story.append(_pdf_paragraph("Prerequisite Knowledge", styles["h2"]))
        _pdf_bullets(story, overview.get("prerequisite_knowledge"), styles)
    if overview.get("core_formulas"):
        story.append(_pdf_paragraph("Core Formula Sheet", styles["h2"]))
        _pdf_table(
            story,
            ["Concept and Formula"],
            [[item] for item in _as_list(overview.get("core_formulas"))],
            [6.7],
            styles,
        )

    objectives = _as_list(result.get("learning_objectives"))
    if objectives:
        story.append(_pdf_paragraph("Learning Objectives", styles["h1"]))
        for index, item in enumerate(objectives, start=1):
            text = item.get("objective") if isinstance(item, dict) else item
            story.append(Paragraph(f"{index}. {escape(_plain_text(text))}", styles["bullet"]))

    sections = _as_list(result.get("lecture_sections"))
    if sections:
        story.append(_pdf_paragraph("Session Plan", styles["h1"]))
        rows = []
        for section in sections:
            if isinstance(section, dict):
                content = section.get("content", {})
                focus = content.get("key_points") if isinstance(content, dict) else content
                rows.append([
                    f"{section.get('minutes', '')} min",
                    section.get("title"),
                    _plain_text(focus)[:250],
                ])
        _pdf_table(story, ["Time", "Section", "Teaching Focus"], rows, [0.75, 2.3, 3.65], styles)
        story.extend([PageBreak(), _pdf_paragraph("Detailed Lecture Notes", styles["h1"])])
        for section in sections:
            if not isinstance(section, dict):
                continue
            story.append(_pdf_paragraph(
                _plain_text(section.get("title"))
                + (f" - {section.get('minutes')} minutes" if section.get("minutes") else ""),
                styles["h2"],
            ))
            _render_teaching_section_pdf(story, section.get("content", {}), styles)

    _render_teaching_assessment_pdf(story, result, styles)
    case_study = result.get("case_study")
    if isinstance(case_study, dict) and case_study:
        story.extend([Spacer(1, 10), _pdf_paragraph("Case Study", styles["h1"])])
        story.append(_pdf_paragraph(case_study.get("title"), styles["h2"]))
        _pdf_callout(story, "Scenario", case_study.get("scenario"), styles, "EEF3FA")
        if case_study.get("questions"):
            story.append(_pdf_paragraph("Case Questions", styles["h3"]))
            for index, question in enumerate(_as_list(case_study.get("questions")), start=1):
                story.append(Paragraph(f"{index}. {escape(_plain_text(question))}", styles["bullet"]))
        _pdf_callout(story, "Teaching notes", case_study.get("teaching_notes"), styles, "F3F0FB")
    _render_sources_pdf(story, result, styles)
    _render_quality_notes_pdf(story, result, styles)


def _render_teaching_section_pdf(story: list, content: Any, styles) -> None:
    if not isinstance(content, dict):
        _pdf_bullets(story, content, styles)
        return
    if content.get("explanation"):
        _pdf_callout(story, "Explanation", content.get("explanation"), styles)
    if content.get("key_points"):
        story.append(_pdf_paragraph("Key Points", styles["h3"]))
        _pdf_bullets(story, content.get("key_points"), styles)
    for label, key, fill in (
        ("Classroom activity", "classroom_activity", "EAF7F8"),
        ("Concept check", "concept_check", "EEF3FA"),
        ("Mini discussion", "mini_discussion", "F3F0FB"),
        ("Teaching tip", "teaching_tip", "FFF7E7"),
        ("Closing prompt", "closing_prompt", "EEF3FA"),
        ("Expected answer", "expected_answer", "EAF6F0"),
    ):
        _pdf_callout(story, label, content.get(key), styles, fill)
    for key in ("worked_example", "worked_examples"):
        for index, example in enumerate(_as_list(content.get(key)), start=1):
            if isinstance(example, dict):
                story.append(_pdf_paragraph(f"Worked Example {index}", styles["h3"]))
                _pdf_callout(
                    story,
                    "Problem",
                    example.get("problem"),
                    styles,
                    "EEF3FA",
                )
                _pdf_callout(
                    story,
                    "Solution",
                    example.get("solution"),
                    styles,
                    "EAF6F0",
                )
    if content.get("board_work"):
        story.append(_pdf_paragraph("Board / Slide Work", styles["h3"]))
        _pdf_bullets(story, content.get("board_work"), styles)
    if content.get("common_errors"):
        story.append(_pdf_paragraph("Common Errors to Address", styles["h3"]))
        _pdf_bullets(story, content.get("common_errors"), styles)
    if content.get("source_refs"):
        story.append(_pdf_paragraph(
            "Source references: " + "; ".join(map(_plain_text, _as_list(content.get("source_refs")))),
            styles["small"],
        ))


def _render_teaching_assessment_pdf(story: list, result: dict, styles) -> None:
    if any(result.get(key) for key in (
        "mcqs", "assignments", "numerical_problems", "bloom_mapping",
        "discussion_questions", "viva_questions"
    )):
        story.extend([PageBreak(), _pdf_paragraph("Assessment and Engagement Toolkit", styles["h1"])])
    mcqs = _as_list(result.get("mcqs"))
    if mcqs:
        story.append(_pdf_paragraph("Multiple-Choice Questions", styles["h2"]))
        for index, item in enumerate(mcqs, start=1):
            if not isinstance(item, dict):
                continue
            story.append(_pdf_paragraph(f"{index}. {_plain_text(item.get('question'))}", styles["h3"]))
            for option_index, option in enumerate(_as_list(item.get("options"))):
                story.append(Paragraph(
                    f"{chr(65 + option_index)}. {escape(_mcq_option_text(option))}",
                    styles["bullet"],
                ))
            _pdf_callout(
                story,
                "Answer and rationale",
                f"{_plain_text(item.get('answer'))}\n{_plain_text(item.get('explanation'))}",
                styles,
                "EAF6F0",
            )
    numerical = _as_list(result.get("numerical_problems"))
    if numerical:
        story.append(_pdf_paragraph("Numerical Problems", styles["h2"]))
        for index, item in enumerate(numerical, start=1):
            if isinstance(item, dict):
                story.append(_pdf_paragraph(f"Problem {index}", styles["h3"]))
                _pdf_callout(story, "Question", item.get("problem"), styles, "EEF3FA")
                _pdf_callout(story, "Given", item.get("given"), styles)
                _pdf_callout(
                    story,
                    "Solution steps",
                    item.get("solution_steps") or item.get("solution"),
                    styles,
                )
                _pdf_callout(story, "Final answer", item.get("final_answer"), styles)
    assignments = _as_list(result.get("assignments"))
    if assignments:
        story.append(_pdf_paragraph("Assignments", styles["h2"]))
        for item in assignments:
            if isinstance(item, dict):
                story.append(_pdf_paragraph(item.get("title"), styles["h3"]))
                _pdf_callout(story, "Brief", item.get("prompt"), styles, "EEF3FA")
                if item.get("deliverables"):
                    story.append(_pdf_paragraph("Deliverables", styles["h3"]))
                    _pdf_bullets(story, item.get("deliverables"), styles)
                rubric_rows = []
                for criterion in _as_list(item.get("rubric")):
                    if isinstance(criterion, dict):
                        rubric_rows.append([
                            criterion.get("criterion"),
                            f"{criterion.get('weight_percent', '')}%",
                            criterion.get("description"),
                        ])
                    else:
                        rubric_rows.append([criterion, "", ""])
                if rubric_rows:
                    _pdf_table(
                        story,
                        ["Criterion", "Weight", "Performance description"],
                        rubric_rows,
                        [1.6, 0.8, 4.3],
                        styles,
                    )
    bloom = _as_list(result.get("bloom_mapping"))
    if bloom:
        rows = [
            [item.get("objective"), item.get("level"), item.get("assessment")]
            for item in bloom
            if isinstance(item, dict)
        ]
        story.append(_pdf_paragraph("Bloom's Taxonomy Mapping", styles["h2"]))
        _pdf_table(
            story,
            ["Learning Objective", "Bloom Level", "Assessment"],
            rows,
            [2.5, 1.2, 3.0],
            styles,
        )
    discussion = _as_list(result.get("discussion_questions"))
    if discussion:
        story.append(_pdf_paragraph("Discussion Questions", styles["h2"]))
        for index, item in enumerate(discussion, start=1):
            story.append(Paragraph(f"{index}. {escape(_plain_text(item))}", styles["bullet"]))
    viva = _as_list(result.get("viva_questions"))
    if viva:
        rows = []
        for item in viva:
            if isinstance(item, dict):
                rows.append([item.get("question"), item.get("expected_answer")])
            else:
                rows.append([item, ""])
        story.append(_pdf_paragraph("Viva Questions and Expected Answers", styles["h2"]))
        _pdf_table(story, ["Question", "Expected Answer / Talking Points"], rows, [3.0, 3.7], styles)


def _render_research_pdf(story: list, result: dict, styles) -> None:
    scope = result.get("scope") if isinstance(result.get("scope"), dict) else {}
    if scope:
        story.append(_pdf_paragraph("Research Scope", styles["h1"]))
        _pdf_callout(story, "Discipline", scope.get("discipline"), styles)
        _pdf_callout(story, "Problem statement", scope.get("problem_statement"), styles)
        _pdf_callout(story, "Included", scope.get("inclusion_boundaries"), styles)
        _pdf_callout(story, "Excluded", scope.get("exclusion_boundaries"), styles)
    story.append(_pdf_paragraph("Executive Summary", styles["h1"]))
    summary = result.get("executive_summary")
    if isinstance(summary, dict):
        _pdf_callout(story, "Evidence", summary.get("evidence"), styles, "EEF3FA")
        _pdf_callout(story, "Interpretation", summary.get("inference"), styles, "F3F0FB")
        _pdf_callout(story, "Limitations", summary.get("limitations"), styles)
    else:
        story.append(_pdf_paragraph(summary, styles["body"]))
    themes = _as_list(result.get("themes"))
    if themes:
        story.append(_pdf_paragraph("Thematic Literature Synthesis", styles["h1"]))
        for index, item in enumerate(themes, start=1):
            if not isinstance(item, dict):
                continue
            block = [_pdf_paragraph(f"{index}. {_plain_text(item.get('theme'))}", styles["h2"])]
            for evidence in _as_list(item.get("evidence")):
                block.append(_pdf_paragraph(f"Evidence: {evidence}", styles["body"]))
            synthesis = item.get("synthesis")
            if isinstance(synthesis, dict):
                block.append(_pdf_paragraph(
                    f"Evidence: {_plain_text(synthesis.get('evidence'))}", styles["body"]
                ))
                block.append(_pdf_paragraph(
                    f"Interpretation: {_plain_text(synthesis.get('inference'))}",
                    styles["body"],
                ))
            if item.get("papers"):
                block.append(_pdf_paragraph(
                    "Studies represented: " + "; ".join(map(_plain_text, _as_list(item.get("papers")))),
                    styles["small"],
                ))
            story.append(KeepTogether(block))
    methods = _as_list(result.get("methodology_comparison"))
    if methods:
        story.extend([PageBreak(), _pdf_paragraph("Methodology Comparison", styles["h1"])])
        for index, item in enumerate(methods, start=1):
            if not isinstance(item, dict):
                continue
            story.append(_pdf_paragraph(f"{index}. {_plain_text(item.get('method'))}", styles["h2"]))
            story.append(_pdf_paragraph(
                f"Strengths: {_plain_text(item.get('strengths'))}", styles["body"]
            ))
            story.append(_pdf_paragraph(
                f"Limitations: {_plain_text(item.get('limitations'))}",
                styles["body"],
            ))
            story.append(_pdf_paragraph(
                f"Suitability: {_plain_text(item.get('suitability'))}", styles["body"]
            ))
            studies = item.get("studies") or item.get("papers")
            if studies:
                story.append(_pdf_paragraph(
                    "Representative studies: "
                    + "; ".join(map(_plain_text, _as_list(studies))),
                    styles["small"],
                ))
    gaps = _as_list(result.get("research_gaps"))
    if gaps:
        story.append(_pdf_paragraph("Research Gaps", styles["h1"]))
        for index, item in enumerate(gaps, start=1):
            if not isinstance(item, dict):
                continue
            story.append(_pdf_paragraph(f"{index}. {_plain_text(item.get('gap'))}", styles["h2"]))
            _pdf_callout(story, "Evidence for the gap", item.get("evidence"), styles, "EEF3FA")
            _pdf_callout(story, "Research implication", item.get("inference"), styles, "F3F0FB")
            _pdf_callout(story, "Confidence", item.get("confidence"), styles)
    questions = _as_list(result.get("research_questions"))
    if questions:
        story.append(_pdf_paragraph("Proposed Research Questions", styles["h1"]))
        for index, question in enumerate(questions, start=1):
            if isinstance(question, dict):
                story.append(Paragraph(
                    f"{index}. {escape(_plain_text(question.get('question')))}", styles["bullet"]
                ))
                _pdf_callout(story, "Rationale", question.get("rationale"), styles)
                _pdf_callout(
                    story,
                    "Key variables or constructs",
                    question.get("key_variables_or_constructs"),
                    styles,
                )
            else:
                story.append(Paragraph(f"{index}. {escape(_plain_text(question))}", styles["bullet"]))
    future = _as_list(result.get("future_scope"))
    if future:
        rows = [[item.get("area"), item.get("scope")] for item in future if isinstance(item, dict)]
        story.append(_pdf_paragraph("Future Scope", styles["h1"]))
        _pdf_table(story, ["Research Area", "Opportunity"], rows, [2.0, 4.7], styles)
    suggestions = _as_list(result.get("methodology_suggestions"))
    if suggestions:
        story.append(_pdf_paragraph("Methodology Recommendations", styles["h1"]))
        for index, item in enumerate(suggestions, start=1):
            if isinstance(item, dict):
                story.append(_pdf_paragraph(
                    f"{index}. {_plain_text(item.get('suggestion'))}", styles["h2"]
                ))
                story.append(_pdf_paragraph(
                    f"Rationale: {_plain_text(item.get('rationale'))}",
                    styles["body"],
                ))
                _pdf_callout(story, "Data collection", item.get("data_collection"), styles)
                _pdf_callout(story, "Analysis", item.get("analysis"), styles)
    references = _as_list(result.get("apa_references"))
    if references:
        story.extend([PageBreak(), _pdf_paragraph("APA References", styles["h1"])])
        for reference in references:
            story.append(_pdf_paragraph(reference, styles["body"]))
    _render_sources_pdf(story, result, styles)
    _render_quality_notes_pdf(story, result, styles)


def _render_sources_pdf(story: list, result: dict, styles) -> None:
    sources = _as_list(result.get("sources"))
    evidence = result.get("overview", {}).get("evidence_base", []) if isinstance(
        result.get("overview"), dict
    ) else []
    if not sources and not evidence:
        return
    story.extend([PageBreak(), _pdf_paragraph("Sources and Further Reading", styles["h1"])])
    for index, item in enumerate(sources, start=1):
        if isinstance(item, dict):
            title = _plain_text(item.get("title"), "Source")
            year = _plain_text(item.get("year"))
            url = _plain_text(item.get("url"))
            story.append(Paragraph(f"{index}. {escape(title)}"
                                   + (f" ({escape(year)})" if year else "")
                                   + (f"<br/><font color='#{MUTED}'>{escape(url)}</font>" if url else ""),
                                   styles["body"]))
        else:
            story.append(Paragraph(f"{index}. {escape(_plain_text(item))}", styles["body"]))
    if evidence:
        story.append(_pdf_paragraph("Evidence Base Named in the Teaching Plan", styles["h2"]))
        _pdf_bullets(story, evidence, styles)


def _render_quality_notes_pdf(story: list, result: dict, styles) -> None:
    notes = result.get("quality_notes")
    if not isinstance(notes, dict):
        return
    story.append(_pdf_paragraph("Professor Review Checklist", styles["h1"]))
    for heading, key in (
        ("Limitations", "limitations"),
        ("Unverified Claims", "unverified_claims"),
        ("Review Priorities", "review_priorities"),
    ):
        values = _as_list(notes.get(key))
        if values:
            story.append(_pdf_paragraph(heading, styles["h2"]))
            _pdf_bullets(story, values, styles)


# ---------------------------------------------------------------------------
# PPTX - professor-ready lecture/research presentation
# ---------------------------------------------------------------------------


def _pptx_set_text(shape, text: Any, *, size: int, color: str, bold: bool = False,
                   align=PP_ALIGN.LEFT):
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    clean_text = _plain_text(text)
    run.text = clean_text
    run.font.name = "Cambria Math" if _contains_math(clean_text) else "Arial"
    run.font.size = PptxPt(size)
    run.font.bold = bold
    run.font.color.rgb = _pptx_color(color)
    return frame


def _pptx_add_text(slide, text: Any, x, y, w, h, *, size=18, color=INK, bold=False,
                   align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(PptxInches(x), PptxInches(y), PptxInches(w), PptxInches(h))
    _pptx_set_text(shape, text, size=size, color=color, bold=bold, align=align)
    return shape


def _pptx_add_header(slide, title: str, section: str, number: int, accent: str):
    clean_title = _shorten(title, 105)
    title_size = 35 if len(clean_title) <= 48 else 30 if len(clean_title) <= 76 else 25
    _pptx_add_text(slide, section.upper(), 0.65, 0.38, 4.6, 0.28, size=10, color=accent, bold=True)
    _pptx_add_text(slide, clean_title, 0.65, 0.72, 11.25, 0.94, size=title_size, color=NAVY, bold=True)
    _pptx_add_text(slide, f"{number:02d}", 12.15, 0.4, 0.5, 0.25, size=9, color=MUTED,
                   align=PP_ALIGN.RIGHT)


def _pptx_add_footer(slide, text: str):
    _pptx_add_text(slide, text, 0.65, 7.16, 12.0, 0.18, size=8, color=MUTED)


def _pptx_add_bullets(
    slide,
    values: Any,
    x,
    y,
    w,
    h,
    *,
    size=19,
    color=INK,
    max_items=7,
    show_bullets=True,
):
    shape = slide.shapes.add_textbox(PptxInches(x), PptxInches(y), PptxInches(w), PptxInches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = PptxInches(0.08)
    frame.margin_right = PptxInches(0.04)
    for index, value in enumerate(_as_list(values)[:max_items]):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        clean_text = _plain_text(value)
        p.text = ("• " if show_bullets else "") + clean_text
        p.level = 0
        p.font.name = "Cambria Math" if _contains_math(clean_text) else "Arial"
        p.font.size = PptxPt(size)
        p.font.color.rgb = _pptx_color(color)
        p.space_after = PptxPt(8)
        p.line_spacing = 1.08
    return shape


def _pptx_add_callout(
    slide,
    label: str,
    text: Any,
    x,
    y,
    w,
    h,
    accent: str,
    *,
    body_size: int = 16,
):
    shape = slide.shapes.add_textbox(PptxInches(x), PptxInches(y), PptxInches(w), PptxInches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = PptxInches(0.16)
    frame.margin_right = PptxInches(0.14)
    frame.margin_top = PptxInches(0.11)
    frame.margin_bottom = PptxInches(0.08)
    p = frame.paragraphs[0]
    r = p.add_run()
    r.text = label.upper()
    r.font.name = "Arial"
    r.font.size = PptxPt(10)
    r.font.bold = True
    r.font.color.rgb = _pptx_color(accent)
    p.space_after = PptxPt(5)
    p2 = frame.add_paragraph()
    clean_text = _plain_text(text)
    p2.text = clean_text
    p2.font.name = "Cambria Math" if _contains_math(clean_text) else "Arial"
    p2.font.size = PptxPt(body_size)
    p2.font.color.rgb = _pptx_color(INK)
    p2.line_spacing = 1.05
    shape.fill.solid()
    shape.fill.fore_color.rgb = _pptx_color(LIGHT)
    shape.line.color.rgb = _pptx_color(accent)
    shape.line.width = PptxPt(1)
    return shape


def _create_pptx(result: dict, path: Path, agent_type: str) -> Path:
    deck = Presentation()
    deck.slide_width = PptxInches(13.333)
    deck.slide_height = PptxInches(7.5)
    accent = CYAN if agent_type == "teaching" else PURPLE
    _pptx_title_slide(deck, result, agent_type, accent)
    if agent_type == "teaching":
        _render_teaching_pptx(deck, result, accent)
    else:
        _render_research_pptx(deck, result, accent)
    deck.core_properties.title = _plain_text(result.get("title"), "Professor AI Output")
    deck.core_properties.subject = "AI-assisted faculty presentation"
    deck.core_properties.author = "Professor AI Workspace"
    deck.save(path)
    return path


def _pptx_title_slide(deck: Presentation, result: dict, agent_type: str, accent: str):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _pptx_color(NAVY)
    _pptx_add_text(
        slide,
        "PROFESSOR AI - "
        + ("TEACHING & CURRICULUM" if agent_type == "teaching" else "RESEARCH SYNTHESIS"),
        0.8, 0.65, 10.8, 0.32, size=12, color=accent, bold=True,
    )
    title = _plain_text(result.get("title"), "Professor AI Output")
    title_size = 42 if len(title) > 80 else 50
    _pptx_add_text(
        slide,
        title,
        0.8, 1.25, 11.5, 3.25, size=title_size, color=WHITE, bold=True,
    )
    overview = result.get("overview") if isinstance(result.get("overview"), dict) else {}
    subtitle = (
        f"{overview.get('course', 'Faculty teaching package')} - "
        f"{overview.get('duration_minutes', '')} minutes - "
        f"{overview.get('difficulty', '')}"
        if agent_type == "teaching"
        else "Literature review, research gaps, questions, and methodology direction"
    )
    _pptx_add_text(slide, subtitle, 0.84, 5.05, 10.9, 0.58, size=20, color="C0CDDD")
    _pptx_add_text(
        slide,
        "AI-assisted faculty draft. Verify facts, citations, calculations, and academic suitability.",
        0.84, 6.72, 11.7, 0.3, size=10, color="8FA4BF",
    )


def _render_teaching_pptx(deck: Presentation, result: dict, accent: str):
    number = 2
    overview = result.get("overview") if isinstance(result.get("overview"), dict) else {}
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _pptx_add_header(slide, "Session at a Glance", "Teaching Package", number, accent)
    number += 1
    _pptx_add_callout(slide, "Lesson purpose", overview.get("lesson_purpose"), 0.7, 1.65, 7.5, 2.0, accent)
    _pptx_add_callout(
        slide,
        "Session design",
        f"Course: {overview.get('course', 'General')}\n"
        f"Audience: {overview.get('audience', 'University students')}\n"
        f"Duration: {overview.get('duration_minutes', '')} minutes\n"
        f"Level: {overview.get('difficulty', 'Intermediate')}",
        8.45, 1.65, 4.2, 2.0, accent,
    )
    _pptx_add_text(slide, "Prerequisite knowledge", 0.75, 4.05, 4.0, 0.35, size=24, color=NAVY, bold=True)
    _pptx_add_bullets(slide, overview.get("prerequisite_knowledge"), 0.75, 4.5, 5.7, 2.05, size=17)
    _pptx_add_text(slide, "Evidence base", 6.9, 4.05, 3.5, 0.35, size=24, color=NAVY, bold=True)
    _pptx_add_bullets(slide, overview.get("evidence_base"), 6.9, 4.5, 5.6, 2.05, size=16)
    _pptx_add_footer(slide, "Professor review edition")

    objectives = _as_list(result.get("learning_objectives"))
    if objectives:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        _pptx_add_header(slide, "Learning Objectives", "Teaching Package", number, accent)
        number += 1
        values = []
        for item in objectives:
            if isinstance(item, dict):
                prefix = f"{item.get('id')}  " if item.get("id") else ""
                values.append(prefix + _plain_text(item.get("objective")))
            else:
                values.append(item)
        _pptx_add_bullets(slide, values, 0.85, 1.65, 11.7, 4.9, size=19, max_items=8)
        _pptx_add_footer(slide, "Learning outcomes should be aligned to assessment and course policy.")

    sections = _as_list(result.get("lecture_sections"))
    if sections:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        _pptx_add_header(slide, "Session Roadmap", "Teaching Package", number, accent)
        number += 1
        y = 1.65
        for section in sections[:7]:
            if not isinstance(section, dict):
                continue
            _pptx_add_text(slide, f"{section.get('minutes', '')} min", 0.8, y, 1.05, 0.34,
                           size=16, color=accent, bold=True)
            _pptx_add_text(slide, section.get("title"), 1.9, y, 10.5, 0.44,
                           size=19, color=NAVY, bold=True)
            y += 0.72
        _pptx_add_footer(slide, "A structured 60-minute teaching flow")

        for section in sections:
            if not isinstance(section, dict):
                continue
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            title = _plain_text(section.get("title"), "Lecture Section")
            _pptx_add_header(slide, title, f"{section.get('minutes', '')} minutes", number, accent)
            number += 1
            content = section.get("content", {})
            if isinstance(content, dict):
                key_points = _as_list(content.get("key_points"))
                _pptx_add_text(slide, "Key teaching points", 0.75, 1.9, 5.8, 0.36,
                               size=24, color=NAVY, bold=True)
                _pptx_add_bullets(slide, key_points, 0.75, 2.33, 6.0, 4.05, size=17, max_items=7)
                activity = (
                    content.get("classroom_activity")
                    or content.get("worked_example")
                    or content.get("worked_examples")
                    or content.get("mini_discussion")
                    or content.get("concept_check")
                )
                _pptx_add_callout(slide, "In-class application", activity, 7.05, 1.92, 5.55, 2.15, accent)
                teaching_note = (
                    content.get("teaching_tip")
                    or content.get("common_errors")
                    or content.get("closing_prompt")
                    or content.get("expected_answer")
                )
                _pptx_add_callout(slide, "Professor note", teaching_note, 7.05, 4.35, 5.55, 1.85, accent)
                refs = _as_list(content.get("source_refs"))
                if refs:
                    _pptx_add_text(slide, "Sources: " + "; ".join(map(_plain_text, refs)),
                                   0.75, 6.55, 11.7, 0.25, size=9, color=MUTED)
            else:
                _pptx_add_bullets(slide, content, 0.85, 1.7, 11.5, 4.8, size=20)

    formulas = _as_list(overview.get("core_formulas"))
    if formulas:
        for start in range(0, len(formulas), 4):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            _pptx_add_header(slide, "Core Formula Sheet", "Reference", number, accent)
            number += 1
            _pptx_add_bullets(slide, formulas[start:start + 4], 0.95, 1.7, 11.2, 4.9, size=22, max_items=4)
            _pptx_add_footer(slide, "Ensure rate and period units are consistent.")

    case_study = result.get("case_study")
    if isinstance(case_study, dict) and case_study:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        _pptx_add_header(slide, _plain_text(case_study.get("title"), "Applied Case"),
                         "Case Study", number, accent)
        number += 1
        _pptx_add_callout(slide, "Scenario", case_study.get("scenario"), 0.75, 1.62, 7.2, 3.5, accent)
        _pptx_add_text(slide, "Discussion questions", 8.25, 1.72, 4.1, 0.4,
                       size=23, color=NAVY, bold=True)
        _pptx_add_bullets(slide, case_study.get("questions"), 8.25, 2.2, 4.15, 3.5, size=17, max_items=5)
        _pptx_add_footer(slide, "Use teaching notes in the professor document for facilitation guidance.")

    discussion = _as_list(result.get("discussion_questions"))
    if discussion:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        _pptx_add_header(slide, "Discussion Questions", "Student Engagement", number, accent)
        number += 1
        _pptx_add_bullets(slide, discussion, 0.9, 1.65, 11.5, 4.9, size=21, max_items=7)

    mcqs = _as_list(result.get("mcqs"))
    for start in range(0, len(mcqs), 3):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        _pptx_add_header(slide, "Knowledge Check", "Assessment", number, accent)
        number += 1
        y = 1.55
        for q_index, item in enumerate(mcqs[start:start + 3], start=start + 1):
            if not isinstance(item, dict):
                continue
            _pptx_add_text(slide, f"{q_index}. {_plain_text(item.get('question'))}",
                           0.75, y, 11.8, 0.6, size=18, color=NAVY, bold=True)
            options = [
                f"{chr(65+i)}. {_mcq_option_text(v)}"
                for i, v in enumerate(_as_list(item.get("options")))
            ]
            _pptx_add_text(slide, "   ".join(options), 0.95, y + 0.58, 11.4, 0.68, size=14, color=INK)
            y += 1.65
        _pptx_add_footer(slide, "Answers and rationales are available in the professor document.")

    assignments = _as_list(result.get("assignments"))
    if assignments:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        _pptx_add_header(slide, "Assignment Brief", "Assessment", number, accent)
        number += 1
        item = assignments[0]
        if isinstance(item, dict):
            _pptx_add_text(slide, item.get("title"), 0.78, 1.62, 11.5, 0.5,
                           size=27, color=NAVY, bold=True)
            _pptx_add_callout(slide, "Task", item.get("prompt"), 0.78, 2.25, 7.5, 2.9, accent)
            _pptx_add_text(slide, "Assessment criteria", 8.55, 2.25, 3.8, 0.4,
                           size=22, color=NAVY, bold=True)
            _pptx_add_bullets(slide, item.get("rubric"), 8.55, 2.75, 3.8, 2.7, size=17, max_items=6)

    _pptx_sources_slide(deck, result, number, accent)


def _render_research_pptx(deck: Presentation, result: dict, accent: str):
    number = 2
    summary = result.get("executive_summary")
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _pptx_add_header(slide, "Executive Summary", "Research Synthesis", number, accent)
    number += 1
    if isinstance(summary, dict):
        _pptx_add_callout(slide, "Evidence", summary.get("evidence"), 0.72, 1.62, 7.6, 4.9, accent)
        _pptx_add_callout(slide, "Interpretation", summary.get("inference"), 8.55, 1.62, 4.05, 4.9, accent)
    else:
        _pptx_add_bullets(slide, summary, 0.85, 1.7, 11.5, 4.8, size=20)

    themes = _as_list(result.get("themes"))
    for item in themes:
        if not isinstance(item, dict):
            continue
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        _pptx_add_header(slide, _plain_text(item.get("theme"), "Research Theme"),
                         "Thematic Synthesis", number, accent)
        number += 1
        synthesis = item.get("synthesis")
        if isinstance(synthesis, dict):
            _pptx_add_callout(slide, "Evidence", synthesis.get("evidence"), 0.72, 1.62, 7.5, 4.45, accent)
            _pptx_add_callout(slide, "Interpretation", synthesis.get("inference"),
                              8.5, 1.62, 4.1, 4.45, accent)
        else:
            _pptx_add_bullets(slide, synthesis, 0.85, 1.7, 11.5, 4.8, size=20)
        papers = _as_list(item.get("papers"))
        if papers:
            _pptx_add_text(slide, "Studies: " + "; ".join(map(_plain_text, papers)),
                           0.75, 6.5, 11.7, 0.27, size=10, color=MUTED)

    gaps = _as_list(result.get("research_gaps"))
    for start in range(0, len(gaps), 2):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        _pptx_add_header(slide, "Research Gaps", "Research Direction", number, accent)
        number += 1
        x_positions = [0.72, 6.82]
        for offset, item in enumerate(gaps[start:start + 2]):
            if not isinstance(item, dict):
                continue
            x = x_positions[offset]
            _pptx_add_text(slide, item.get("gap"), x, 1.62, 5.75, 0.85,
                           size=23, color=NAVY, bold=True)
            _pptx_add_callout(slide, "Evidence", item.get("evidence"), x, 2.58, 5.75, 1.75, accent)
            _pptx_add_callout(slide, "Research implication", item.get("inference"),
                              x, 4.52, 5.75, 1.55, accent)

    questions = _as_list(result.get("research_questions"))
    if questions:
        for start in range(0, len(questions), 6):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            _pptx_add_header(slide, "Proposed Research Questions", "Research Direction", number, accent)
            number += 1
            _pptx_add_bullets(slide, questions[start:start + 6], 0.85, 1.65, 11.5, 4.95,
                              size=19, max_items=6)

    methods = _as_list(result.get("methodology_suggestions"))
    if methods:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        _pptx_add_header(slide, "Methodology Recommendations", "Research Design", number, accent)
        number += 1
        y = 1.55
        for index, item in enumerate(methods[:4], start=1):
            if isinstance(item, dict):
                _pptx_add_text(slide, f"{index}. {_plain_text(item.get('suggestion'))}",
                               0.75, y, 5.4, 0.72, size=19, color=NAVY, bold=True)
                _pptx_add_text(slide, item.get("rationale"), 6.35, y, 6.0, 0.92,
                               size=15, color=INK)
                y += 1.25

    future = _as_list(result.get("future_scope"))
    if future:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        _pptx_add_header(slide, "Future Research Scope", "Next Horizon", number, accent)
        number += 1
        y = 1.55
        for item in future[:5]:
            if isinstance(item, dict):
                _pptx_add_text(slide, item.get("area"), 0.75, y, 3.3, 0.55,
                               size=19, color=accent, bold=True)
                _pptx_add_text(slide, item.get("scope"), 4.25, y, 8.0, 0.75,
                               size=16, color=INK)
                y += 1.0

    _pptx_sources_slide(deck, result, number, accent)


def _pptx_sources_slide(deck: Presentation, result: dict, number: int, accent: str):
    sources = _as_list(result.get("sources"))
    references = _as_list(result.get("apa_references"))
    values = references or [
        (
            _plain_text(item.get("title"))
            + (f" ({_plain_text(item.get('year'))})" if item.get("year") else "")
            + (f" - {_plain_text(item.get('url'))}" if item.get("url") else "")
        )
        if isinstance(item, dict)
        else _plain_text(item)
        for item in sources
    ]
    if not values:
        return
    for start in range(0, len(values), 7):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        _pptx_add_header(slide, "Sources and Further Reading", "References", number, accent)
        number += 1
        _pptx_add_bullets(slide, values[start:start + 7], 0.82, 1.55, 11.75, 5.25,
                          size=15, max_items=7)
        _pptx_add_footer(slide, "Verify citation details before academic publication or distribution.")


# ---------------------------------------------------------------------------
# PPTX v2 layouts: complete, paginated, and readable at presentation distance.
# These definitions intentionally replace the earlier compact renderers above.
# ---------------------------------------------------------------------------


def _pptx_item_text(value: Any) -> str:
    if isinstance(value, dict):
        return "; ".join(
            f"{_humanize(key)}: {_plain_text(item)}"
            for key, item in value.items()
            if item not in (None, "", [])
        )
    return _plain_text(value)


def _pptx_pages(values: Any, *, max_items: int = 6, max_chars: int = 760) -> list[list[str]]:
    pages: list[list[str]] = []
    current: list[str] = []
    current_chars = 0
    for value in _as_list(values):
        text = _pptx_item_text(value)
        if not text:
            continue
        if current and (len(current) >= max_items or current_chars + len(text) > max_chars):
            pages.append(current)
            current = []
            current_chars = 0
        current.append(text)
        current_chars += len(text)
    if current:
        pages.append(current)
    return pages


def _pptx_text_pages(value: Any, *, max_chars: int = 900) -> list[str]:
    text = _plain_text(value)
    if not text:
        return []
    paragraphs = [part.strip() for part in text.split("\n") if part.strip()]
    pages: list[str] = []
    current = ""
    for paragraph in paragraphs or [text]:
        words = paragraph.split()
        while words:
            available = max_chars - len(current)
            piece: list[str] = []
            while words and len(" ".join(piece + [words[0]])) <= max(80, available):
                piece.append(words.pop(0))
            if not piece:
                piece.append(words.pop(0))
            candidate = (current + "\n\n" + " ".join(piece)).strip()
            if current and len(candidate) > max_chars:
                pages.append(current)
                current = " ".join(piece)
            else:
                current = candidate
    if current:
        pages.append(current)
    return pages


def _pptx_new_slide(deck, title: str, section: str, number: int, accent: str):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _pptx_color("FBFCFE")
    _pptx_add_header(slide, title, section, number, accent)
    return slide


def _pptx_list_slides(
    deck,
    *,
    title: str,
    section: str,
    values: Any,
    number: int,
    accent: str,
    size: int = 19,
    footer: str = "",
) -> int:
    pages = _pptx_pages(values)
    for page_index, page in enumerate(pages, start=1):
        page_title = title if len(pages) == 1 else f"{title} ({page_index} of {len(pages)})"
        slide = _pptx_new_slide(deck, page_title, section, number, accent)
        number += 1
        _pptx_add_bullets(slide, page, 0.85, 1.65, 11.65, 4.95, size=size, max_items=len(page))
        if footer:
            _pptx_add_footer(slide, footer)
    return number


def _pptx_prose_slides(
    deck,
    *,
    title: str,
    section: str,
    value: Any,
    number: int,
    accent: str,
    label: str = "",
) -> int:
    pages = _pptx_text_pages(value)
    for page_index, page in enumerate(pages, start=1):
        page_title = title if len(pages) == 1 else f"{title} ({page_index} of {len(pages)})"
        slide = _pptx_new_slide(deck, page_title, section, number, accent)
        number += 1
        if label:
            _pptx_add_text(slide, label.upper(), 0.85, 1.72, 11.4, 0.3, size=11, color=accent, bold=True)
        _pptx_add_text(slide, page, 0.85, 2.12 if label else 1.72, 11.55, 4.55,
                       size=18, color=INK)
    return number


def _render_teaching_pptx(deck: Presentation, result: dict, accent: str):
    number = 2
    overview = result.get("overview") if isinstance(result.get("overview"), dict) else {}
    slide = _pptx_new_slide(deck, "Session Overview", "Teaching Package", number, accent)
    number += 1
    _pptx_add_text(slide, overview.get("lesson_purpose"), 0.8, 1.6, 7.4, 1.45,
                   size=23, color=NAVY, bold=True)
    session_details = [
        f"Course: {overview.get('course', 'General')}",
        f"Audience: {overview.get('audience', 'University students')}",
        f"Duration: {overview.get('duration_minutes', '')} minutes",
        f"Difficulty: {overview.get('difficulty', 'intermediate')}",
    ]
    _pptx_add_bullets(slide, session_details, 8.55, 1.62, 3.8, 2.25, size=17, max_items=4)
    _pptx_add_text(slide, "Prerequisite knowledge", 0.8, 3.65, 5.6, 0.4, size=22, color=NAVY, bold=True)
    _pptx_add_bullets(slide, overview.get("prerequisite_knowledge"), 0.8, 4.14, 5.65, 2.25,
                      size=17, max_items=6)
    _pptx_add_text(slide, "Evidence base", 6.8, 3.65, 5.5, 0.4, size=22, color=NAVY, bold=True)
    _pptx_add_bullets(slide, overview.get("evidence_base"), 6.8, 4.14, 5.55, 2.25,
                      size=16, max_items=6)

    objectives = []
    for item in _as_list(result.get("learning_objectives")):
        if isinstance(item, dict):
            objectives.append(
                f"{item.get('id')}: {item.get('objective')} [{item.get('bloom_level')}]"
            )
        else:
            objectives.append(item)
    number = _pptx_list_slides(
        deck, title="Learning Objectives", section="Teaching Package", values=objectives,
        number=number, accent=accent, footer="Each objective is mapped to an assessment.",
    )

    sections = [item for item in _as_list(result.get("lecture_sections")) if isinstance(item, dict)]
    roadmap = [f"{item.get('minutes', '')} min - {item.get('title', 'Lecture section')}" for item in sections]
    number = _pptx_list_slides(
        deck, title="Session Roadmap", section="Teaching Package", values=roadmap,
        number=number, accent=accent,
    )
    for section_item in sections:
        section_title = _plain_text(section_item.get("title"), "Lecture Section")
        content = section_item.get("content") if isinstance(section_item.get("content"), dict) else {}
        number = _pptx_prose_slides(
            deck, title=section_title, section=f"{section_item.get('minutes', '')} minutes",
            value=content.get("explanation"), number=number, accent=accent, label="Explanation",
        )
        number = _pptx_list_slides(
            deck, title=f"{section_title}: Key Points", section="Lecture Notes",
            values=content.get("key_points"), number=number, accent=accent,
        )
        application = []
        for label, key in (
            ("Worked example", "worked_example"),
            ("Classroom activity", "classroom_activity"),
            ("Concept check", "concept_check"),
            ("Teaching tip", "teaching_tip"),
        ):
            if content.get(key):
                application.append(f"{label}: {_plain_text(content.get(key))}")
        application.extend(f"Common error: {_plain_text(item)}" for item in _as_list(content.get("common_errors")))
        number = _pptx_list_slides(
            deck, title=f"{section_title}: Application", section="Lecture Notes",
            values=application, number=number, accent=accent, size=17,
        )

    number = _pptx_list_slides(
        deck, title="Core Formula Sheet", section="Reference",
        values=overview.get("core_formulas"), number=number, accent=accent, size=20,
    )

    case_study = result.get("case_study")
    if isinstance(case_study, dict):
        number = _pptx_prose_slides(
            deck, title=_plain_text(case_study.get("title"), "Applied Case"),
            section="Case Study", value=case_study.get("scenario"), number=number,
            accent=accent, label="Scenario",
        )
        number = _pptx_list_slides(
            deck, title="Case Questions", section="Case Study", values=case_study.get("questions"),
            number=number, accent=accent,
            footer="Facilitation notes are included in the professor handbook.",
        )

    number = _pptx_list_slides(
        deck, title="Discussion Questions", section="Student Engagement",
        values=result.get("discussion_questions"), number=number, accent=accent,
    )

    for index, item in enumerate(_as_list(result.get("mcqs")), start=1):
        if not isinstance(item, dict):
            continue
        slide = _pptx_new_slide(deck, f"Knowledge Check {index}", "Assessment", number, accent)
        number += 1
        _pptx_add_text(slide, item.get("question"), 0.8, 1.5, 11.7, 0.9,
                       size=22, color=NAVY, bold=True)
        options = [
            f"{chr(65 + i)}. {_mcq_option_text(option)}"
            for i, option in enumerate(_as_list(item.get("options")))
        ]
        _pptx_add_bullets(slide, options, 0.95, 2.55, 11.4, 2.75, size=18, max_items=4)
        answer = f"Answer: {item.get('answer')} - {_plain_text(item.get('explanation'))}"
        _pptx_add_text(slide, answer, 0.85, 5.65, 11.5, 0.78, size=15, color=MUTED)

    for item in _as_list(result.get("assignments")):
        if not isinstance(item, dict):
            continue
        values = [f"Task: {_plain_text(item.get('prompt'))}"]
        values.extend(f"Deliverable: {_plain_text(value)}" for value in _as_list(item.get("deliverables")))
        for criterion in _as_list(item.get("rubric")):
            if isinstance(criterion, dict):
                values.append(
                    f"{criterion.get('criterion')} ({criterion.get('weight_percent')}%): "
                    f"{criterion.get('description')}"
                )
        number = _pptx_list_slides(
            deck, title=_plain_text(item.get("title"), "Assignment"), section="Assignment",
            values=values, number=number, accent=accent, size=17,
        )

    for index, item in enumerate(_as_list(result.get("numerical_problems")), start=1):
        if not isinstance(item, dict):
            continue
        values = [f"Problem: {_plain_text(item.get('problem'))}"]
        values.extend(f"Given: {_plain_text(value)}" for value in _as_list(item.get("given")))
        values.extend(
            f"Step {step_index}: {_plain_text(value)}"
            for step_index, value in enumerate(_as_list(item.get("solution_steps")), start=1)
        )
        values.append(f"Final answer: {_plain_text(item.get('final_answer'))}")
        number = _pptx_list_slides(
            deck, title=f"Numerical Problem {index}", section="Assessment",
            values=values, number=number, accent=accent, size=17,
        )

    number = _pptx_list_slides(
        deck, title="Bloom's Taxonomy Mapping", section="Assessment",
        values=result.get("bloom_mapping"), number=number, accent=accent, size=17,
    )
    number = _pptx_list_slides(
        deck, title="Viva Questions", section="Assessment",
        values=result.get("viva_questions"), number=number, accent=accent, size=17,
    )
    number = _pptx_quality_slides(deck, result, number, accent)
    _pptx_sources_slide(deck, result, number, accent)


def _render_research_pptx(deck: Presentation, result: dict, accent: str):
    number = 2
    scope = result.get("scope") if isinstance(result.get("scope"), dict) else {}
    scope_values = [
        f"Discipline: {_plain_text(scope.get('discipline'))}",
        f"Problem statement: {_plain_text(scope.get('problem_statement'))}",
    ]
    scope_values.extend(f"Included: {_plain_text(value)}" for value in _as_list(scope.get("inclusion_boundaries")))
    scope_values.extend(f"Excluded: {_plain_text(value)}" for value in _as_list(scope.get("exclusion_boundaries")))
    number = _pptx_list_slides(
        deck, title="Research Scope", section="Research Synthesis", values=scope_values,
        number=number, accent=accent, size=18,
    )

    summary = result.get("executive_summary")
    if isinstance(summary, dict):
        for label, key in (("Evidence", "evidence"), ("Interpretation", "inference"), ("Limitations", "limitations")):
            number = _pptx_prose_slides(
                deck, title=f"Executive Summary: {label}", section="Research Synthesis",
                value=summary.get(key), number=number, accent=accent, label=label,
            )

    for item in _as_list(result.get("themes")):
        if not isinstance(item, dict):
            continue
        values = [f"Evidence: {_plain_text(value)}" for value in _as_list(item.get("evidence"))]
        synthesis = item.get("synthesis") if isinstance(item.get("synthesis"), dict) else {}
        values.extend([
            f"Synthesis: {_plain_text(synthesis.get('evidence'))}",
            f"Interpretation: {_plain_text(synthesis.get('inference'))}",
        ])
        if item.get("papers"):
            values.append("Studies: " + "; ".join(map(_plain_text, _as_list(item.get("papers")))))
        number = _pptx_list_slides(
            deck, title=_plain_text(item.get("theme"), "Research Theme"),
            section="Thematic Synthesis", values=values, number=number, accent=accent, size=17,
        )

    for item in _as_list(result.get("methodology_comparison")):
        if not isinstance(item, dict):
            continue
        values = []
        values.extend(f"Study: {_plain_text(value)}" for value in _as_list(item.get("studies")))
        values.extend(f"Strength: {_plain_text(value)}" for value in _as_list(item.get("strengths")))
        values.extend(f"Limitation: {_plain_text(value)}" for value in _as_list(item.get("limitations")))
        values.append(f"Suitability: {_plain_text(item.get('suitability'))}")
        number = _pptx_list_slides(
            deck, title=_plain_text(item.get("method"), "Method"),
            section="Methodology Comparison", values=values, number=number, accent=accent, size=17,
        )

    for item in _as_list(result.get("research_gaps")):
        if not isinstance(item, dict):
            continue
        values = [
            f"Evidence: {_plain_text(item.get('evidence'))}",
            f"Research implication: {_plain_text(item.get('inference'))}",
            f"Confidence: {_plain_text(item.get('confidence'))}",
        ]
        number = _pptx_list_slides(
            deck, title=_plain_text(item.get("gap"), "Research Gap"), section="Research Gaps",
            values=values, number=number, accent=accent, size=18,
        )

    for item in _as_list(result.get("research_questions")):
        if not isinstance(item, dict):
            continue
        values = [
            f"Question: {_plain_text(item.get('question'))}",
            f"Rationale: {_plain_text(item.get('rationale'))}",
            "Key variables or constructs: "
            + "; ".join(map(_plain_text, _as_list(item.get("key_variables_or_constructs")))),
        ]
        number = _pptx_list_slides(
            deck, title="Proposed Research Question", section="Research Direction",
            values=values, number=number, accent=accent, size=18,
        )

    number = _pptx_list_slides(
        deck, title="Future Research Scope", section="Research Direction",
        values=result.get("future_scope"), number=number, accent=accent, size=18,
    )
    for item in _as_list(result.get("methodology_suggestions")):
        if not isinstance(item, dict):
            continue
        values = [
            f"Rationale: {_plain_text(item.get('rationale'))}",
            f"Data collection: {_plain_text(item.get('data_collection'))}",
            f"Analysis: {_plain_text(item.get('analysis'))}",
        ]
        number = _pptx_list_slides(
            deck, title=_plain_text(item.get("suggestion"), "Methodology Recommendation"),
            section="Research Design", values=values, number=number, accent=accent, size=18,
        )
    number = _pptx_quality_slides(deck, result, number, accent)
    _pptx_sources_slide(deck, result, number, accent)


def _pptx_quality_slides(deck, result: dict, number: int, accent: str) -> int:
    notes = result.get("quality_notes")
    if not isinstance(notes, dict):
        return number
    values = []
    values.extend(f"Limitation: {_plain_text(value)}" for value in _as_list(notes.get("limitations")))
    values.extend(
        f"Unverified claim: {_plain_text(value)}" for value in _as_list(notes.get("unverified_claims"))
    )
    values.extend(
        f"Review priority: {_plain_text(value)}" for value in _as_list(notes.get("review_priorities"))
    )
    return _pptx_list_slides(
        deck, title="Professor Review Checklist", section="Quality Review", values=values,
        number=number, accent=accent, size=17,
    )


# ---------------------------------------------------------------------------
# PPTX v3 layouts: a concise classroom/research deck, not a paginated handbook.
# The PDF and DOCX retain the complete prose; slides surface the material a presenter needs.
# ---------------------------------------------------------------------------


def _pptx_complete_excerpt(value: Any, limit: int) -> str:
    """Fit slide copy at a complete sentence boundary instead of cutting mid-sentence."""
    text = _pptx_item_text(value).strip()
    text = re.sub(r"(?<=[.!?]);\s+(?=[A-Z][A-Za-z /-]{1,32}:)", " ", text)
    text = re.sub(r";\s+(?=[A-Z][A-Za-z /-]{1,32}:)", ". ", text)
    if len(text) <= limit:
        return text
    sentences = re.split(r"(?<=[.!?])\s+", text)
    selected: list[str] = []
    used = 0
    for sentence in sentences:
        sentence = sentence.strip()
        if not sentence:
            continue
        candidate_length = used + len(sentence) + (1 if selected else 0)
        if selected and candidate_length > limit:
            break
        selected.append(sentence)
        used = candidate_length
        if used >= limit:
            break
    # A single long sentence is safer than a clipped fragment; callout auto-fit handles it.
    return " ".join(selected or sentences[:1]).strip()


def _pptx_estimated_lines(text: Any, width: float, font_size: int) -> int:
    clean_text = _plain_text(text)
    if not clean_text:
        return 0
    usable_width = max(0.5, width - 0.34)
    average_character_width = max(0.07, font_size * 0.52 / 72)
    characters_per_line = max(12, int(usable_width / average_character_width))
    return sum(
        max(1, (len(line.strip()) + characters_per_line - 1) // characters_per_line)
        for line in clean_text.splitlines() or [clean_text]
    )


def _pptx_callout_height(
    text: Any,
    width: float,
    *,
    body_size: int = 16,
    minimum: float = 0.82,
    maximum: float = 3.5,
) -> float:
    body_lines = _pptx_estimated_lines(text, width, body_size)
    natural = 0.42 + body_lines * (body_size * 1.13 / 72)
    return min(maximum, max(minimum, natural))


def _pptx_compact_values(values: Any, *, item_limit: int = 150) -> list[str]:
    return [
        _pptx_complete_excerpt(value, item_limit)
        for value in _as_list(values)
        if _plain_text(value)
    ]


def _pptx_compact_list_slides(
    deck,
    *,
    title: str,
    section: str,
    values: Any,
    number: int,
    accent: str,
    max_items: int = 6,
    item_limit: int = 150,
    size: int = 19,
    footer: str = "",
) -> int:
    compact = _pptx_compact_values(values, item_limit=item_limit)
    pages = [compact[index:index + max_items] for index in range(0, len(compact), max_items)]
    for page_index, page in enumerate(pages, start=1):
        page_title = title if len(pages) == 1 else f"{title} ({page_index} of {len(pages)})"
        slide = _pptx_new_slide(deck, page_title, section, number, accent)
        number += 1
        _pptx_add_bullets(slide, page, 0.82, 1.82, 11.7, 4.95, size=size, max_items=len(page))
        if footer:
            _pptx_add_footer(slide, footer)
    return number


def _pptx_section_slide(deck, section_item: dict, number: int, accent: str) -> int:
    title = _plain_text(section_item.get("title"), "Lecture Section")
    content = section_item.get("content") if isinstance(section_item.get("content"), dict) else {}
    slide = _pptx_new_slide(
        deck, title, f"{section_item.get('minutes', '')} minutes", number, accent
    )
    explanation = _pptx_complete_excerpt(content.get("explanation"), 320)
    if explanation:
        _pptx_add_text(slide, explanation, 0.76, 1.72, 11.75, 0.82, size=16, color=MUTED)
    _pptx_add_text(slide, "Key teaching points", 0.76, 2.68, 5.7, 0.35,
                   size=22, color=NAVY, bold=True)
    points = _pptx_compact_values(content.get("key_points"), item_limit=90)[:4]
    _pptx_add_bullets(slide, points, 0.76, 3.12, 5.85, 3.0, size=16, max_items=4)

    worked = content.get("worked_example") or content.get("worked_examples")
    application = _pptx_complete_excerpt(
        worked or content.get("classroom_activity") or content.get("concept_check"), 430
    )
    professor_note = _pptx_complete_excerpt(
        content.get("teaching_tip") or content.get("common_errors"), 320
    )
    right_x = 6.9
    right_width = 5.7
    top = 2.62
    bottom = 6.35
    gap = 0.18
    if application and professor_note:
        application_height = _pptx_callout_height(
            application, right_width, body_size=15, minimum=1.35, maximum=2.35
        )
        note_height = _pptx_callout_height(
            professor_note, right_width, body_size=15, minimum=1.05, maximum=1.85
        )
        available = bottom - top - gap
        total = application_height + note_height
        if total > available:
            scale = available / total
            application_height *= scale
            note_height *= scale
        _pptx_add_callout(
            slide,
            "Application",
            application,
            right_x,
            top,
            right_width,
            application_height,
            accent,
            body_size=15,
        )
        _pptx_add_callout(
            slide,
            "Professor note",
            professor_note,
            right_x,
            top + application_height + gap,
            right_width,
            note_height,
            accent,
            body_size=15,
        )
    else:
        single_text = application or professor_note
        if single_text:
            label = "Application" if application else "Professor note"
            height = _pptx_callout_height(
                single_text, right_width, body_size=15, minimum=1.05, maximum=bottom - top
            )
            _pptx_add_callout(
                slide, label, single_text, right_x, top, right_width, height, accent,
                body_size=15,
            )
    refs = _as_list(content.get("source_refs"))
    if refs:
        _pptx_add_footer(slide, "Sources: " + "; ".join(map(_plain_text, refs)))
    return number + 1


def _pptx_mcq_slide(deck, item: dict, index: int, number: int, accent: str) -> int:
    slide = _pptx_new_slide(deck, f"Knowledge Check {index}", "Assessment", number, accent)
    question = _pptx_complete_excerpt(item.get("question"), 260)
    question_height = min(
        1.25,
        max(0.62, _pptx_estimated_lines(question, 11.75, 20) * 0.34),
    )
    _pptx_add_text(slide, question, 0.78, 1.65, 11.75, question_height,
                   size=20, color=NAVY, bold=True)
    options = [
        f"{chr(65 + option_index)}. "
        f"{_pptx_complete_excerpt(_mcq_option_text(option), 220)}"
        for option_index, option in enumerate(_as_list(item.get("options"))[:4])
    ]
    answer = (
        f"Answer: {item.get('answer')}  "
        f"{_pptx_complete_excerpt(item.get('explanation'), 380)}"
    )
    answer_height = _pptx_callout_height(
        answer, 11.75, body_size=15, minimum=0.92, maximum=1.45
    )
    answer_y = 6.72 - answer_height
    options_y = 1.65 + question_height + 0.18
    options_height = max(1.65, answer_y - options_y - 0.25)
    _pptx_add_bullets(
        slide,
        options,
        0.88,
        options_y,
        11.45,
        options_height,
        size=17,
        max_items=4,
        show_bullets=False,
    )
    _pptx_add_callout(
        slide,
        "Answer and rationale",
        answer,
        0.78,
        answer_y,
        11.75,
        answer_height,
        accent,
        body_size=15,
    )
    return number + 1


def _render_teaching_pptx(deck: Presentation, result: dict, accent: str):
    number = 2
    overview = result.get("overview") if isinstance(result.get("overview"), dict) else {}
    slide = _pptx_new_slide(deck, "Session Overview", "Teaching Package", number, accent)
    number += 1
    _pptx_add_callout(slide, "Lesson purpose", _pptx_complete_excerpt(overview.get("lesson_purpose"), 420),
                      0.72, 1.75, 7.45, 1.82, accent)
    details = [
        f"Course: {_plain_text(overview.get('course'), 'General')}",
        f"Audience: {_plain_text(overview.get('audience'), 'University students')}",
        f"Duration: {overview.get('duration_minutes', '')} minutes",
        f"Difficulty: {_plain_text(overview.get('difficulty'), 'intermediate')}",
    ]
    _pptx_add_bullets(slide, details, 8.55, 1.8, 3.75, 1.9, size=17, max_items=4)
    _pptx_add_text(slide, "Prerequisite knowledge", 0.76, 4.02, 5.7, 0.35,
                   size=22, color=NAVY, bold=True)
    _pptx_add_bullets(slide, _pptx_compact_values(overview.get("prerequisite_knowledge"), item_limit=80)[:3],
                      0.76, 4.45, 5.7, 1.75, size=16, max_items=3)
    _pptx_add_text(slide, "Evidence base", 6.82, 4.02, 5.5, 0.35,
                   size=22, color=NAVY, bold=True)
    _pptx_add_bullets(slide, _pptx_compact_values(overview.get("evidence_base"), item_limit=85)[:3],
                      6.82, 4.45, 5.55, 1.75, size=15, max_items=3)

    objectives = [
        f"{item.get('id')}: {_pptx_complete_excerpt(item.get('objective'), 145)} "
        f"[{item.get('bloom_level')}]"
        if isinstance(item, dict) else _pptx_complete_excerpt(item, 160)
        for item in _as_list(result.get("learning_objectives"))
    ]
    number = _pptx_compact_list_slides(
        deck, title="Learning Objectives", section="Teaching Package", values=objectives,
        number=number, accent=accent, max_items=6, size=18,
        footer="Each objective is mapped to an assessment.",
    )
    sections = [item for item in _as_list(result.get("lecture_sections")) if isinstance(item, dict)]
    roadmap = [
        f"{item.get('minutes', '')} min - {_pptx_complete_excerpt(item.get('title'), 115)}"
        for item in sections
    ]
    number = _pptx_compact_list_slides(
        deck, title="Session Roadmap", section="Teaching Package", values=roadmap,
        number=number, accent=accent, max_items=7, item_limit=140, size=18,
    )
    for section_item in sections:
        number = _pptx_section_slide(deck, section_item, number, accent)

    number = _pptx_compact_list_slides(
        deck, title="Core Formula Sheet", section="Reference",
        values=overview.get("core_formulas"), number=number, accent=accent,
        max_items=6, item_limit=155, size=18,
    )

    case_study = result.get("case_study")
    if isinstance(case_study, dict) and case_study:
        slide = _pptx_new_slide(
            deck, _plain_text(case_study.get("title"), "Applied Case"), "Case Study", number, accent
        )
        number += 1
        _pptx_add_callout(slide, "Scenario", _pptx_complete_excerpt(case_study.get("scenario"), 560),
                          0.72, 1.78, 7.35, 4.35, accent)
        _pptx_add_text(slide, "Case questions", 8.35, 1.82, 4.0, 0.4,
                       size=22, color=NAVY, bold=True)
        _pptx_add_bullets(
            slide, _pptx_compact_values(case_study.get("questions"), item_limit=125)[:5],
            8.35, 2.35, 4.0, 3.55, size=16, max_items=5,
        )

    number = _pptx_compact_list_slides(
        deck, title="Discussion Questions", section="Student Engagement",
        values=result.get("discussion_questions"), number=number, accent=accent,
        max_items=6, item_limit=150, size=18,
    )
    for index, item in enumerate(_as_list(result.get("mcqs")), start=1):
        if isinstance(item, dict):
            number = _pptx_mcq_slide(deck, item, index, number, accent)

    for item in _as_list(result.get("assignments")):
        if not isinstance(item, dict):
            continue
        slide = _pptx_new_slide(deck, _plain_text(item.get("title"), "Assignment"),
                                "Assignment", number, accent)
        number += 1
        _pptx_add_callout(slide, "Task", _pptx_complete_excerpt(item.get("prompt"), 520),
                          0.72, 1.78, 7.35, 3.9, accent)
        criteria = []
        for criterion in _as_list(item.get("rubric")):
            if isinstance(criterion, dict):
                criteria.append(
                    f"{criterion.get('criterion')} ({criterion.get('weight_percent')}%): "
                    f"{_pptx_complete_excerpt(criterion.get('description'), 95)}"
                )
        _pptx_add_text(slide, "Assessment criteria", 8.35, 1.82, 4.0, 0.4,
                       size=22, color=NAVY, bold=True)
        _pptx_add_bullets(slide, criteria[:5], 8.35, 2.35, 4.0, 3.5, size=16, max_items=5)

    for index, item in enumerate(_as_list(result.get("numerical_problems")), start=1):
        if not isinstance(item, dict):
            continue
        slide = _pptx_new_slide(deck, f"Numerical Problem {index}", "Assessment", number, accent)
        number += 1
        problem_text = _pptx_complete_excerpt(item.get("problem"), 420)
        given_text = _pptx_complete_excerpt(item.get("given"), 320)
        prompt = problem_text + (f"\n\nGiven: {given_text}" if given_text else "")
        problem_height = _pptx_callout_height(
            prompt, 11.85, body_size=15, minimum=1.05, maximum=2.35
        )
        _pptx_add_callout(
            slide,
            "Problem",
            prompt,
            0.72,
            1.72,
            11.85,
            problem_height,
            accent,
            body_size=15,
        )
        steps = [f"{i}. {_pptx_complete_excerpt(step, 185)}" for i, step in enumerate(
            _as_list(item.get("solution_steps"))[:6], start=1
        )]
        solution_y = 1.72 + problem_height + 0.28
        _pptx_add_text(slide, "Solution", 0.76, solution_y, 7.35, 0.35,
                       size=21, color=NAVY, bold=True)
        answer = _pptx_complete_excerpt(item.get("final_answer"), 280)
        answer_height = _pptx_callout_height(
            answer, 4.1, body_size=16, minimum=0.88, maximum=1.45
        )
        body_y = solution_y + 0.43
        body_height = max(1.55, 6.45 - body_y)
        _pptx_add_bullets(
            slide,
            steps,
            0.76,
            body_y,
            7.35,
            body_height,
            size=16,
            max_items=6,
            show_bullets=False,
        )
        _pptx_add_callout(
            slide,
            "Final answer",
            answer,
            8.45,
            solution_y,
            4.1,
            answer_height,
            accent,
        )

    number = _pptx_compact_list_slides(
        deck, title="Bloom's Taxonomy Mapping", section="Assessment",
        values=result.get("bloom_mapping"), number=number, accent=accent,
        max_items=6, item_limit=145, size=17,
    )
    number = _pptx_compact_list_slides(
        deck, title="Viva Questions", section="Assessment", values=result.get("viva_questions"),
        number=number, accent=accent, max_items=5, item_limit=180, size=17,
    )
    number = _pptx_quality_slides(deck, result, number, accent)
    _pptx_sources_slide(deck, result, number, accent)


def _render_research_pptx(deck: Presentation, result: dict, accent: str):
    number = 2
    scope = result.get("scope") if isinstance(result.get("scope"), dict) else {}
    scope_values = [
        f"Discipline: {_plain_text(scope.get('discipline'))}",
        f"Problem statement: {_pptx_complete_excerpt(scope.get('problem_statement'), 260)}",
        *[
            f"Included: {_pptx_complete_excerpt(value, 150)}"
            for value in _as_list(scope.get("inclusion_boundaries"))
        ],
        *[
            f"Excluded: {_pptx_complete_excerpt(value, 150)}"
            for value in _as_list(scope.get("exclusion_boundaries"))
        ],
    ]
    number = _pptx_compact_list_slides(
        deck, title="Research Scope", section="Research Synthesis", values=scope_values,
        number=number, accent=accent, max_items=6, item_limit=280, size=18,
    )

    summary = result.get("executive_summary")
    if isinstance(summary, dict):
        slide = _pptx_new_slide(deck, "Executive Summary", "Research Synthesis", number, accent)
        number += 1
        _pptx_add_callout(slide, "Evidence", _pptx_complete_excerpt(summary.get("evidence"), 500),
                          0.72, 1.78, 7.35, 2.1, accent)
        _pptx_add_callout(slide, "Interpretation", _pptx_complete_excerpt(summary.get("inference"), 360),
                          8.35, 1.78, 4.2, 2.1, accent)
        _pptx_add_callout(slide, "Limitations", _pptx_complete_excerpt(summary.get("limitations"), 520),
                          0.72, 4.2, 11.83, 1.65, accent)

    for item in _as_list(result.get("themes")):
        if not isinstance(item, dict):
            continue
        synthesis = item.get("synthesis") if isinstance(item.get("synthesis"), dict) else {}
        slide = _pptx_new_slide(deck, _plain_text(item.get("theme"), "Research Theme"),
                                "Thematic Synthesis", number, accent)
        number += 1
        evidence = item.get("evidence") or synthesis.get("evidence")
        _pptx_add_callout(slide, "Evidence", _pptx_complete_excerpt(evidence, 560),
                          0.72, 1.78, 7.35, 3.9, accent)
        _pptx_add_callout(slide, "Interpretation", _pptx_complete_excerpt(synthesis.get("inference"), 420),
                          8.35, 1.78, 4.2, 3.9, accent)

    number = _pptx_compact_list_slides(
        deck, title="Methodology Comparison", section="Research Design",
        values=result.get("methodology_comparison"), number=number, accent=accent,
        max_items=4, item_limit=260, size=17,
    )
    for item in _as_list(result.get("research_gaps")):
        if not isinstance(item, dict):
            continue
        slide = _pptx_new_slide(deck, _plain_text(item.get("gap"), "Research Gap"),
                                "Research Gaps", number, accent)
        number += 1
        _pptx_add_callout(slide, "Evidence", _pptx_complete_excerpt(item.get("evidence"), 520),
                          0.72, 1.78, 7.35, 3.7, accent)
        implication = (
            f"{_pptx_complete_excerpt(item.get('inference'), 340)}\n\nConfidence: "
            f"{_plain_text(item.get('confidence'))}"
        )
        _pptx_add_callout(slide, "Research implication", implication,
                          8.35, 1.78, 4.2, 3.7, accent)

    number = _pptx_compact_list_slides(
        deck, title="Proposed Research Questions", section="Research Direction",
        values=result.get("research_questions"), number=number, accent=accent,
        max_items=4, item_limit=230, size=17,
    )
    number = _pptx_compact_list_slides(
        deck, title="Future Research Scope", section="Research Direction",
        values=result.get("future_scope"), number=number, accent=accent,
        max_items=5, item_limit=210, size=17,
    )
    number = _pptx_compact_list_slides(
        deck, title="Methodology Recommendations", section="Research Design",
        values=result.get("methodology_suggestions"), number=number, accent=accent,
        max_items=4, item_limit=260, size=17,
    )
    number = _pptx_quality_slides(deck, result, number, accent)
    _pptx_sources_slide(deck, result, number, accent)


def _pptx_quality_slides(deck, result: dict, number: int, accent: str) -> int:
    notes = result.get("quality_notes")
    if not isinstance(notes, dict):
        return number
    values = [
        *[
            f"Limitation: {_pptx_complete_excerpt(value, 150)}"
            for value in _as_list(notes.get("limitations"))
        ],
        *[
            f"Unverified claim: {_pptx_complete_excerpt(value, 150)}"
            for value in _as_list(notes.get("unverified_claims"))
        ],
        *[
            f"Review priority: {_pptx_complete_excerpt(value, 150)}"
            for value in _as_list(notes.get("review_priorities"))
        ],
    ]
    return _pptx_compact_list_slides(
        deck, title="Professor Review Checklist", section="Quality Review", values=values,
        number=number, accent=accent, max_items=6, item_limit=170, size=17,
    )


def _pptx_sources_slide(deck: Presentation, result: dict, number: int, accent: str):
    references = _as_list(result.get("apa_references"))
    sources = _as_list(result.get("sources"))
    values = references or [
        (
            _plain_text(item.get("title"), "Source")
            + (f" ({_plain_text(item.get('year'))})" if item.get("year") else "")
            + (f" - {_clean_url(_plain_text(item.get('url')))}" if item.get("url") else "")
        )
        if isinstance(item, dict) else _plain_text(item)
        for item in sources
    ]
    _pptx_compact_list_slides(
        deck, title="Sources and Further Reading", section="References", values=values,
        number=number, accent=accent, max_items=5, item_limit=210, size=15,
        footer="Verify citation details before academic publication or distribution.",
    )
